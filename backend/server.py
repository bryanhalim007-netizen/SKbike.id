from dotenv import load_dotenv
from pathlib import Path
import os

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

from fastapi import FastAPI, APIRouter, HTTPException, Request, Response, Depends, UploadFile, File, Header, Query
from starlette.middleware.cors import CORSMiddleware
from starlette.responses import Response as StarletteResponse
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel, Field, EmailStr, ConfigDict, BeforeValidator
from typing import List, Optional, Annotated
from bson import ObjectId
import logging
import uuid
import io
import csv
import json
import html
import re
import bcrypt
import jwt
import requests
from datetime import datetime, timezone, timedelta

# ---------------- DB ----------------
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

app = FastAPI()
api_router = APIRouter(prefix="/api")

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

# ---------------- Auth utils ----------------
JWT_ALGORITHM = "HS256"
WHATSAPP_NUMBER = os.environ.get("WHATSAPP_NUMBER", "628125559681")

def get_jwt_secret() -> str:
    return os.environ["JWT_SECRET"]

def hash_password(password: str) -> str:
    return bcrypt.hashpw(password.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")

def verify_password(plain: str, hashed: str) -> bool:
    return bcrypt.checkpw(plain.encode("utf-8"), hashed.encode("utf-8"))

def create_access_token(user_id: str, email: str) -> str:
    payload = {"sub": user_id, "email": email, "exp": datetime.now(timezone.utc) + timedelta(minutes=60), "type": "access"}
    return jwt.encode(payload, get_jwt_secret(), algorithm=JWT_ALGORITHM)

def create_refresh_token(user_id: str) -> str:
    payload = {"sub": user_id, "exp": datetime.now(timezone.utc) + timedelta(days=7), "type": "refresh"}
    return jwt.encode(payload, get_jwt_secret(), algorithm=JWT_ALGORITHM)

def set_auth_cookies(response: Response, access_token: str, refresh_token: str):
    response.set_cookie(key="access_token", value=access_token, httponly=True, secure=True, samesite="none", max_age=3600, path="/")
    response.set_cookie(key="refresh_token", value=refresh_token, httponly=True, secure=True, samesite="none", max_age=604800, path="/")

async def get_current_user(request: Request) -> dict:
    token = request.cookies.get("access_token")
    if not token:
        auth_header = request.headers.get("Authorization", "")
        if auth_header.startswith("Bearer "):
            token = auth_header[7:]
    if not token:
        raise HTTPException(status_code=401, detail="Belum terautentikasi")
    try:
        payload = jwt.decode(token, get_jwt_secret(), algorithms=[JWT_ALGORITHM])
        if payload.get("type") != "access":
            raise HTTPException(status_code=401, detail="Tipe token tidak valid")
        user = await db.users.find_one({"_id": ObjectId(payload["sub"])})
        if not user:
            raise HTTPException(status_code=401, detail="Pengguna tidak ditemukan")
        user["_id"] = str(user["_id"])
        user.pop("password_hash", None)
        return user
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Token kedaluwarsa")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Token tidak valid")

def get_super_admin_email() -> str:
    return os.environ["ADMIN_EMAIL"].lower()

def is_super_admin(user: dict) -> bool:
    return (user.get("email") or "").lower() == get_super_admin_email()

def get_role(user: dict) -> str:
    if is_super_admin(user):
        return "super_admin"
    return (user.get("role") or "admin").lower()

def can_see_financials(user: dict) -> bool:
    # Only the super admin sees margin & omset figures.
    return is_super_admin(user)

async def forbid_cashier(user: dict = Depends(get_current_user)) -> dict:
    if get_role(user) == "cashier":
        raise HTTPException(status_code=403, detail="Akses ditolak: akun kasir tidak memiliki izin ini")
    return user

def strip_sale_financials(pub: dict, user: dict) -> dict:
    if not can_see_financials(user):
        pub["margin"] = None
        pub["harga_modal"] = None
    return pub

async def require_super_admin(user: dict = Depends(get_current_user)) -> dict:
    if not is_super_admin(user):
        raise HTTPException(status_code=403, detail="Hanya admin utama yang dapat mengelola akun admin")
    return user

# ---------------- Activity log ----------------
async def log_activity(user: dict, action: str, entity: str, label: str, entity_id: str = None):
    """Record who (admin/kasir) created/changed/deleted a record."""
    try:
        await db.activity_logs.insert_one({
            "id": str(uuid.uuid4()),
            "actor_email": user.get("email", ""),
            "actor_name": user.get("name", "") or user.get("email", ""),
            "actor_role": get_role(user),
            "action": action,
            "entity": entity,
            "entity_id": entity_id,
            "label": label,
            "created_at": datetime.now(timezone.utc).isoformat(),
        })
    except Exception:
        pass

def activity_public(d: dict) -> dict:
    return {
        "id": d.get("id"),
        "actor_email": d.get("actor_email"),
        "actor_name": d.get("actor_name"),
        "actor_role": d.get("actor_role"),
        "action": d.get("action"),
        "entity": d.get("entity"),
        "entity_id": d.get("entity_id"),
        "label": d.get("label"),
        "created_at": d.get("created_at"),
    }

# ---------------- Brute force protection ----------------
MAX_FAILED_ATTEMPTS = 5
LOCKOUT_MINUTES = 15

async def check_lockout(identifier: str):
    rec = await db.login_attempts.find_one({"identifier": identifier})
    if rec and rec.get("count", 0) >= MAX_FAILED_ATTEMPTS:
        locked_until = rec.get("locked_until")
        if locked_until:
            lu = datetime.fromisoformat(locked_until)
            if lu > datetime.now(timezone.utc):
                remaining = int((lu - datetime.now(timezone.utc)).total_seconds() // 60) + 1
                raise HTTPException(status_code=429, detail=f"Terlalu banyak percobaan gagal. Coba lagi dalam {remaining} menit.")

async def register_failed_attempt(identifier: str):
    rec = await db.login_attempts.find_one({"identifier": identifier})
    count = (rec.get("count", 0) if rec else 0) + 1
    update = {"count": count, "last_attempt": datetime.now(timezone.utc).isoformat()}
    if count >= MAX_FAILED_ATTEMPTS:
        update["locked_until"] = (datetime.now(timezone.utc) + timedelta(minutes=LOCKOUT_MINUTES)).isoformat()
    await db.login_attempts.update_one({"identifier": identifier}, {"$set": update}, upsert=True)

async def clear_failed_attempts(identifier: str):
    await db.login_attempts.delete_one({"identifier": identifier})

# ---------------- Object storage ----------------
STORAGE_BASE = (os.environ.get("INTEGRATION_PROXY_URL") or "").strip() or "https://integrations.emergentagent.com"
STORAGE_URL = STORAGE_BASE.rstrip("/") + "/objstore/api/v1/storage"
EMERGENT_KEY = os.environ.get("EMERGENT_LLM_KEY")
APP_NAME = "velox-bike"
storage_key = None

MIME_TYPES = {"jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png", "gif": "image/gif", "webp": "image/webp"}

def init_storage(force: bool = False):
    global storage_key
    if storage_key and not force:
        return storage_key
    resp = requests.post(f"{STORAGE_URL}/init", json={"emergent_key": EMERGENT_KEY}, timeout=30)
    resp.raise_for_status()
    storage_key = resp.json()["storage_key"]
    return storage_key

def put_object(path: str, data: bytes, content_type: str) -> dict:
    key = init_storage()
    resp = requests.put(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key, "Content-Type": content_type}, data=data, timeout=120)
    if resp.status_code == 404:
        key = init_storage(force=True)
        resp = requests.put(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key, "Content-Type": content_type}, data=data, timeout=120)
    resp.raise_for_status()
    return resp.json()

def get_object(path: str):
    key = init_storage()
    resp = requests.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=60)
    if resp.status_code == 404:
        key = init_storage(force=True)
        resp = requests.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=60)
    resp.raise_for_status()
    return resp.content, resp.headers.get("Content-Type", "application/octet-stream")

# ---------------- Models ----------------
CATEGORIES = ["Sepeda Listrik", "Sepeda Gunung", "BMX", "Sepeda Anak", "Sepeda Lipat", "Motor / Mobil Aki", "Mini Trail", "Road Bike"]

class Specs(BaseModel):
    frame: Optional[str] = ""
    transmisi: Optional[str] = ""
    rem: Optional[str] = ""
    ukuran_roda: Optional[str] = ""
    baterai_motor: Optional[str] = ""

class Variant(BaseModel):
    color: str
    stock: int = 0
    hex: str = "#94A3B8"
    image_url: str = ""

class Size(BaseModel):
    label: str
    code: str = ""
    cost_price: float = 0
    price: float = 0
    colors: Optional[List[Variant]] = None

class ProductCreate(BaseModel):
    name: str
    code: str = ""
    category: str
    description: str = ""
    price: float = 0
    cost_price: float = 0
    stock: int = 0
    status: str = "Tersedia"
    image_url: str = ""
    supplier: str = ""
    specs: Specs = Field(default_factory=Specs)
    variants: Optional[List[Variant]] = None
    sizes: Optional[List[Size]] = None

class ProductUpdate(BaseModel):
    name: Optional[str] = None
    code: Optional[str] = None
    category: Optional[str] = None
    description: Optional[str] = None
    price: Optional[float] = None
    cost_price: Optional[float] = None
    stock: Optional[int] = None
    status: Optional[str] = None
    image_url: Optional[str] = None
    supplier: Optional[str] = None
    specs: Optional[Specs] = None
    variants: Optional[List[Variant]] = None
    sizes: Optional[List[Size]] = None

def normalize_variants(variants) -> list:
    out = []
    for v in variants or []:
        color = (v.color if hasattr(v, "color") else v.get("color", "")).strip()
        if not color:
            continue
        raw = v.stock if hasattr(v, "stock") else v.get("stock", 0)
        hex_val = (v.hex if hasattr(v, "hex") else v.get("hex", "")) or "#94A3B8"
        img = (v.image_url if hasattr(v, "image_url") else v.get("image_url", "")) or ""
        out.append({"color": color, "stock": max(0, int(raw or 0)), "hex": str(hex_val), "image_url": str(img)})
    return out

def normalize_sizes(sizes) -> list:
    out = []
    for s in sizes or []:
        label = (s.label if hasattr(s, "label") else s.get("label", "")).strip()
        if not label:
            continue
        code = (s.code if hasattr(s, "code") else s.get("code", "")) or ""
        cost = s.cost_price if hasattr(s, "cost_price") else s.get("cost_price", 0)
        price = s.price if hasattr(s, "price") else s.get("price", 0)
        raw_colors = s.colors if hasattr(s, "colors") else s.get("colors")
        colors = normalize_variants(raw_colors)
        if not colors:
            colors = [{"color": "Standar", "stock": 0, "hex": "#94A3B8"}]
        out.append({
            "label": label,
            "code": str(code).strip(),
            "cost_price": float(cost or 0),
            "price": float(price or 0),
            "colors": colors,
        })
    return out

def aggregate_variants_from_sizes(sizes) -> list:
    # Flatten nested size->color stock into a single color list (compat layer
    # for catalog chips, POS, and legacy stock views).
    agg = {}
    order = []
    for s in sizes:
        for c in s.get("colors", []):
            key = c["color"]
            if key not in agg:
                agg[key] = {"color": key, "hex": c.get("hex", "#94A3B8"), "stock": 0, "image_url": c.get("image_url", "") or ""}
                order.append(key)
            agg[key]["stock"] += int(c.get("stock", 0) or 0)
            if not agg[key].get("image_url") and c.get("image_url"):
                agg[key]["image_url"] = c.get("image_url")
    return [agg[k] for k in order]

def product_public(doc: dict) -> dict:
    variants = doc.get("variants") or []
    stock = sum(int(v.get("stock", 0)) for v in variants) if variants else int(doc.get("stock", 0) or 0)
    # Katalog publik TANPA harga: buang harga jual & modal dari tiap ukuran.
    public_sizes = [{k: v for k, v in s.items() if k not in ("price", "cost_price")} for s in (doc.get("sizes", []) or [])]
    return {
        "id": str(doc["_id"]),
        "name": doc.get("name"),
        "category": doc.get("category"),
        "description": doc.get("description", ""),
        "stock": stock,
        "status": doc.get("status", "Tersedia"),
        "image_url": doc.get("image_url", ""),
        "specs": doc.get("specs", {}),
        "variants": variants,
        "sizes": public_sizes,
        "featured": bool(doc.get("featured", False)),
        "sort_order": doc.get("sort_order", 0),
    }

def product_admin(doc: dict, user: dict = None) -> dict:
    pub = product_public(doc)
    pub["price"] = doc.get("price", 0)
    pub["code"] = doc.get("code", "")
    pub["supplier"] = doc.get("supplier", "")
    full_sizes = doc.get("sizes", []) or []
    show_cost = user is None or is_super_admin(user)
    # Modal (cost price) is visible only to the super admin.
    if show_cost:
        pub["cost_price"] = doc.get("cost_price", 0)
        pub["sizes"] = full_sizes
    else:
        pub["sizes"] = [{k: v for k, v in s.items() if k != "cost_price"} for s in full_sizes]
    return pub

# ---------------- Kasir (Kalkulator) models ----------------
class SaleCreate(BaseModel):
    tanggal_penjualan: Optional[str] = None
    nama_pembeli: Optional[str] = None
    nama_barang: Optional[str] = None
    kode_barang: Optional[str] = None
    ukuran_warna: Optional[str] = None
    kode_huruf: Optional[str] = None
    harga_modal: Optional[float] = None
    harga_jual: Optional[float] = None
    margin: Optional[float] = None
    metode_pembayaran: Optional[str] = None
    sudah_diambil: Optional[str] = None
    metode_pengambilan: Optional[str] = None
    alamat_pengiriman: Optional[str] = None
    foto_produk: Optional[str] = None
    bukti_transfer: Optional[str] = None

SALE_FIELDS = [
    "tanggal_penjualan", "nama_pembeli", "nama_barang", "kode_barang", "ukuran_warna",
    "kode_huruf", "harga_modal", "harga_jual", "margin", "metode_pembayaran",
    "sudah_diambil", "metode_pengambilan", "alamat_pengiriman", "foto_produk", "bukti_transfer",
]

def sale_public(doc: dict) -> dict:
    out = {"id": doc.get("id") or str(doc.get("_id")), "created_at": doc.get("created_at"), "cashier": doc.get("cashier", ""), "owner_email": doc.get("cashier", ""), "source": doc.get("source", "kasir"), "service_id": doc.get("service_id"), "transaction_id": doc.get("transaction_id"), "qty": doc.get("qty"), "returned": bool(doc.get("returned", False)), "returned_at": doc.get("returned_at")}
    for f in SALE_FIELDS:
        out[f] = doc.get(f)
    return out

# ---------------- Service (Bengkel) models ----------------
SERVICE_STATUSES = ["Antri", "Dikerjakan", "Menunggu Sparepart", "Selesai", "Diambil"]

class ServiceCreate(BaseModel):
    nomor_servis: Optional[str] = None
    nama_customer: Optional[str] = None
    no_hp: Optional[str] = None
    sepeda: Optional[str] = None
    keluhan: Optional[str] = None
    kondisi_awal: Optional[str] = None
    teknisi: Optional[str] = None
    sparepart: Optional[str] = None
    harga_sparepart: Optional[float] = 0
    biaya_jasa: Optional[float] = 0
    estimasi_selesai: Optional[str] = None
    status: Optional[str] = "Antri"

SERVICE_FIELDS = [
    "nomor_servis", "nama_customer", "no_hp", "sepeda", "keluhan", "kondisi_awal", "teknisi",
    "sparepart", "harga_sparepart", "biaya_jasa", "estimasi_selesai", "status",
]

def service_public(doc: dict) -> dict:
    out = {"id": doc.get("id") or str(doc.get("_id")), "created_at": doc.get("created_at"), "owner_email": doc.get("owner_email", "")}
    for f in SERVICE_FIELDS:
        out[f] = doc.get(f)
    return out

# ---------------- POS (Point of Sale) models ----------------
BULAN_ID = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus", "September", "Oktober", "November", "Desember"]

def id_date_label(dt: datetime) -> str:
    return f"{dt.day:02d} {BULAN_ID[dt.month - 1]} {dt.year}"

def rupiah_id(n) -> str:
    return "Rp " + f"{int(round(float(n or 0))):,}".replace(",", ".")

class PosItem(BaseModel):
    product_id: Optional[str] = None
    name: str
    code: Optional[str] = ""
    color: Optional[str] = None
    size: Optional[str] = None
    qty: int = 1
    unit_price: float = 0
    cost_price: Optional[float] = 0

class PosCheckout(BaseModel):
    customer: Optional[str] = None
    customer_phone: Optional[str] = None
    address: Optional[str] = None
    taken_status: Optional[str] = None
    courier_name: Optional[str] = None
    courier_phone: Optional[str] = None
    transfer_proof: Optional[str] = None
    product_note: Optional[str] = None
    payment_method: Optional[str] = "Cash"
    amount_paid: Optional[float] = 0
    discount: Optional[float] = 0
    items: List[PosItem] = []

# ---------------- Auth routes ----------------
class LoginInput(BaseModel):
    email: EmailStr
    password: str

@api_router.post("/auth/login")
async def login(payload: LoginInput, request: Request, response: Response):
    email = payload.email.lower()
    fwd = request.headers.get("X-Forwarded-For", "")
    ip = (fwd.split(",")[0].strip() if fwd else "") or request.headers.get("X-Real-IP", "") or (request.client.host if request.client else "unknown")
    identifier = f"{ip}:{email}"
    await check_lockout(identifier)
    user = await db.users.find_one({"email": email})
    if not user or not verify_password(payload.password, user["password_hash"]):
        await register_failed_attempt(identifier)
        raise HTTPException(status_code=401, detail="Email atau kata sandi salah")
    await clear_failed_attempts(identifier)
    uid = str(user["_id"])
    access = create_access_token(uid, email)
    refresh = create_refresh_token(uid)
    set_auth_cookies(response, access, refresh)
    return {"id": uid, "email": email, "name": user.get("name", "Admin"), "role": ("super_admin" if email == get_super_admin_email() else user.get("role", "admin")), "is_super": email == get_super_admin_email(), "token": access}

@api_router.post("/auth/refresh")
async def refresh_token(request: Request, response: Response):
    token = request.cookies.get("refresh_token")
    if not token:
        raise HTTPException(status_code=401, detail="Tidak ada token refresh")
    try:
        payload = jwt.decode(token, get_jwt_secret(), algorithms=[JWT_ALGORITHM])
        if payload.get("type") != "refresh":
            raise HTTPException(status_code=401, detail="Tipe token tidak valid")
        user = await db.users.find_one({"_id": ObjectId(payload["sub"])})
        if not user:
            raise HTTPException(status_code=401, detail="Pengguna tidak ditemukan")
        uid = str(user["_id"])
        access = create_access_token(uid, user["email"])
        response.set_cookie(key="access_token", value=access, httponly=True, secure=True, samesite="none", max_age=3600, path="/")
        return {"id": uid, "email": user["email"], "name": user.get("name", "Admin"), "role": get_role(user), "is_super": is_super_admin(user), "token": access}
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Sesi berakhir, silakan masuk lagi")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Token tidak valid")

@api_router.post("/auth/logout")
async def logout(response: Response):
    response.delete_cookie("access_token", path="/")
    response.delete_cookie("refresh_token", path="/")
    return {"message": "Berhasil keluar"}

@api_router.get("/auth/me")
async def me(user: dict = Depends(get_current_user)):
    return {"id": user["_id"], "email": user["email"], "name": user.get("name", "Admin"), "role": get_role(user), "is_super": is_super_admin(user)}

# ---------------- Admin account management (super admin only) ----------------
class AdminCreate(BaseModel):
    email: EmailStr
    password: str
    name: str = "Admin"
    role: str = "admin"

class PasswordChange(BaseModel):
    password: str

class RoleChange(BaseModel):
    role: str

def admin_public(u: dict) -> dict:
    email = (u.get("email") or "").lower()
    is_super = email == get_super_admin_email()
    return {
        "id": str(u.get("_id")),
        "email": u.get("email"),
        "name": u.get("name", "Admin"),
        "role": "super_admin" if is_super else (u.get("role") or "admin"),
        "is_super": is_super,
        "created_at": u.get("created_at"),
    }

@api_router.get("/admin/admins")
async def list_admins(user: dict = Depends(require_super_admin)):
    docs = await db.users.find({}).sort("created_at", 1).to_list(200)
    return [admin_public(d) for d in docs]

@api_router.post("/admin/admins")
async def create_admin(payload: AdminCreate, user: dict = Depends(require_super_admin)):
    email = payload.email.lower()
    if len(payload.password) < 6:
        raise HTTPException(status_code=400, detail="Kata sandi minimal 6 karakter")
    if await db.users.find_one({"email": email}):
        raise HTTPException(status_code=400, detail="Email sudah terdaftar")
    role = (payload.role or "admin").lower()
    if role not in ("admin", "cashier"):
        raise HTTPException(status_code=400, detail="Role tidak valid")
    doc = {"email": email, "password_hash": hash_password(payload.password), "name": payload.name or "Admin", "role": role, "created_at": datetime.now(timezone.utc).isoformat()}
    res = await db.users.insert_one(doc)
    doc["_id"] = res.inserted_id
    await log_activity(user, "create", "account", f"Buat akun {email} (role: {role})", str(res.inserted_id))
    return admin_public(doc)

@api_router.delete("/admin/admins/{admin_id}")
async def delete_admin(admin_id: str, user: dict = Depends(require_super_admin)):
    try:
        oid = ObjectId(admin_id)
    except Exception:
        raise HTTPException(status_code=400, detail="ID tidak valid")
    target = await db.users.find_one({"_id": oid})
    if not target:
        raise HTTPException(status_code=404, detail="Admin tidak ditemukan")
    if (target.get("email") or "").lower() == get_super_admin_email():
        raise HTTPException(status_code=400, detail="Admin utama tidak dapat dihapus")
    if str(target["_id"]) == user["_id"]:
        raise HTTPException(status_code=400, detail="Tidak dapat menghapus akun sendiri")
    await db.users.delete_one({"_id": oid})
    await log_activity(user, "delete", "account", f"Hapus akun {target.get('email')}", admin_id)
    return {"ok": True}

@api_router.put("/admin/admins/{admin_id}/password")
async def change_admin_password(admin_id: str, payload: PasswordChange, user: dict = Depends(require_super_admin)):
    if len(payload.password) < 6:
        raise HTTPException(status_code=400, detail="Kata sandi minimal 6 karakter")
    try:
        oid = ObjectId(admin_id)
    except Exception:
        raise HTTPException(status_code=400, detail="ID tidak valid")
    target = await db.users.find_one({"_id": oid})
    if not target:
        raise HTTPException(status_code=404, detail="Admin tidak ditemukan")
    await db.users.update_one({"_id": oid}, {"$set": {"password_hash": hash_password(payload.password)}})
    return {"ok": True}

@api_router.put("/admin/admins/{admin_id}/role")
async def change_admin_role(admin_id: str, payload: RoleChange, user: dict = Depends(require_super_admin)):
    role = (payload.role or "").lower()
    if role not in ("admin", "cashier"):
        raise HTTPException(status_code=400, detail="Role tidak valid")
    try:
        oid = ObjectId(admin_id)
    except Exception:
        raise HTTPException(status_code=400, detail="ID tidak valid")
    target = await db.users.find_one({"_id": oid})
    if not target:
        raise HTTPException(status_code=404, detail="Admin tidak ditemukan")
    if (target.get("email") or "").lower() == get_super_admin_email():
        raise HTTPException(status_code=400, detail="Role admin utama tidak dapat diubah")
    await db.users.update_one({"_id": oid}, {"$set": {"role": role}})
    updated = await db.users.find_one({"_id": oid})
    await log_activity(user, "update", "account", f"Ubah role {target.get('email')} → {role}", admin_id)
    return admin_public(updated)

@api_router.get("/admin/activity")
async def list_activity(limit: int = 150, user: dict = Depends(require_super_admin)):
    docs = await db.activity_logs.find({}).sort("created_at", -1).to_list(min(max(int(limit or 150), 1), 500))
    return [activity_public(d) for d in docs]

# ---------------- Access PIN (super admin manages, any admin verifies) ----------------
DEFAULT_PINS = {"produk": "1614", "kasir": "1515"}

async def get_access_pins() -> dict:
    doc = await db.settings.find_one({"key": "access_pins"})
    if not doc:
        return dict(DEFAULT_PINS)
    return {"produk": doc.get("produk", DEFAULT_PINS["produk"]), "kasir": doc.get("kasir", DEFAULT_PINS["kasir"])}

class PinsUpdate(BaseModel):
    produk: Optional[str] = None
    kasir: Optional[str] = None

class PinVerify(BaseModel):
    scope: str
    pin: str

@api_router.get("/admin/pins")
async def get_pins(user: dict = Depends(require_super_admin)):
    return await get_access_pins()

@api_router.put("/admin/pins")
async def update_pins(payload: PinsUpdate, user: dict = Depends(require_super_admin)):
    update = {}
    for field in ("produk", "kasir"):
        val = getattr(payload, field)
        if val is not None:
            if not (val.isdigit() and len(val) == 4):
                raise HTTPException(status_code=400, detail="PIN harus 4 angka")
            update[field] = val
    if not update:
        raise HTTPException(status_code=400, detail="Tidak ada PIN untuk diperbarui")
    await db.settings.update_one({"key": "access_pins"}, {"$set": {"key": "access_pins", **update}}, upsert=True)
    return await get_access_pins()

@api_router.post("/admin/verify-pin")
async def verify_pin(payload: PinVerify, user: dict = Depends(get_current_user)):
    pins = await get_access_pins()
    if payload.scope not in pins:
        raise HTTPException(status_code=400, detail="Scope tidak valid")
    return {"ok": pins[payload.scope] == payload.pin}

# ---------------- Absensi / Pegawai ----------------
ATTENDANCE_STATUSES = {"Hadir", "Izin", "Sakit", "Alpa"}

def wib_today() -> str:
    return (datetime.now(timezone.utc) + timedelta(hours=7)).strftime("%Y-%m-%d")

def _valid_date(s: str) -> bool:
    try:
        datetime.strptime(s, "%Y-%m-%d")
        return True
    except (ValueError, TypeError):
        return False

class EmployeeCreate(BaseModel):
    nama: str
    umur: Optional[int] = None
    alamat: Optional[str] = None
    no_hp: Optional[str] = None
    jabatan: Optional[str] = None
    cv_url: Optional[str] = None
    ktp_url: Optional[str] = None
    foto_url: Optional[str] = None

class EmployeeUpdate(BaseModel):
    nama: Optional[str] = None
    umur: Optional[int] = None
    alamat: Optional[str] = None
    no_hp: Optional[str] = None
    jabatan: Optional[str] = None
    cv_url: Optional[str] = None
    ktp_url: Optional[str] = None
    foto_url: Optional[str] = None

class AttendanceMark(BaseModel):
    employee_id: str
    date: str
    status: str

def employee_public(e: dict) -> dict:
    return {
        "id": e.get("id"),
        "nama": e.get("nama"),
        "umur": e.get("umur"),
        "alamat": e.get("alamat"),
        "no_hp": e.get("no_hp"),
        "jabatan": e.get("jabatan"),
        "cv_url": e.get("cv_url"),
        "ktp_url": e.get("ktp_url"),
        "foto_url": e.get("foto_url"),
        "owner_id": e.get("owner_id"),
        "owner_email": e.get("owner_email"),
        "owner_name": e.get("owner_name"),
        "created_at": e.get("created_at"),
    }

async def _get_employee_or_403(emp_id: str, user: dict, require_owner: bool = False) -> dict:
    emp = await db.employees.find_one({"id": emp_id})
    if not emp:
        raise HTTPException(status_code=404, detail="Pegawai tidak ditemukan")
    is_owner = emp.get("owner_id") == user["_id"]
    if require_owner and not is_owner:
        raise HTTPException(status_code=403, detail="Tidak dapat mengubah data pegawai admin lain")
    if not is_owner and not is_super_admin(user):
        raise HTTPException(status_code=403, detail="Tidak diizinkan")
    return emp

@api_router.get("/admin/employees")
async def list_employees(user: dict = Depends(get_current_user)):
    query = {} if is_super_admin(user) else {"owner_id": user["_id"]}
    docs = await db.employees.find(query).sort("created_at", 1).to_list(2000)
    today = wib_today()
    result = []
    for d in docs:
        att = await db.attendance.find_one({"employee_id": d["id"], "date": today})
        pub = employee_public(d)
        pub["today_status"] = att.get("status") if att else None
        result.append(pub)
    return result

@api_router.post("/admin/employees")
async def create_employee(payload: EmployeeCreate, user: dict = Depends(get_current_user)):
    if not payload.nama.strip():
        raise HTTPException(status_code=400, detail="Nama pegawai wajib diisi")
    doc = {
        "id": str(uuid.uuid4()),
        "owner_id": user["_id"],
        "owner_email": user["email"],
        "owner_name": user.get("name", "Admin"),
        "nama": payload.nama.strip(),
        "umur": payload.umur,
        "alamat": (payload.alamat or "").strip() or None,
        "no_hp": (payload.no_hp or "").strip() or None,
        "jabatan": (payload.jabatan or "").strip() or None,
        "cv_url": (payload.cv_url or "").strip() or None,
        "ktp_url": (payload.ktp_url or "").strip() or None,
        "foto_url": (payload.foto_url or "").strip() or None,
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.employees.insert_one(doc)
    pub = employee_public(doc)
    pub["today_status"] = None
    return pub

@api_router.get("/admin/employees/{emp_id}")
async def get_employee(emp_id: str, user: dict = Depends(get_current_user)):
    emp = await _get_employee_or_403(emp_id, user)
    return employee_public(emp)

@api_router.put("/admin/employees/{emp_id}")
async def update_employee(emp_id: str, payload: EmployeeUpdate, user: dict = Depends(get_current_user)):
    # Owner or super admin may edit employee details
    emp = await _get_employee_or_403(emp_id, user)
    updates = {}
    if payload.nama is not None:
        if not payload.nama.strip():
            raise HTTPException(status_code=400, detail="Nama pegawai wajib diisi")
        updates["nama"] = payload.nama.strip()
    if payload.umur is not None:
        updates["umur"] = payload.umur
    if payload.alamat is not None:
        updates["alamat"] = payload.alamat.strip() or None
    if payload.no_hp is not None:
        updates["no_hp"] = payload.no_hp.strip() or None
    if payload.jabatan is not None:
        updates["jabatan"] = payload.jabatan.strip() or None
    if payload.cv_url is not None:
        updates["cv_url"] = payload.cv_url.strip() or None
    if payload.ktp_url is not None:
        updates["ktp_url"] = payload.ktp_url.strip() or None
    if payload.foto_url is not None:
        updates["foto_url"] = payload.foto_url.strip() or None
    if updates:
        updates["updated_at"] = datetime.now(timezone.utc).isoformat()
        await db.employees.update_one({"id": emp_id}, {"$set": updates})
    emp = await db.employees.find_one({"id": emp_id})
    today = wib_today()
    att = await db.attendance.find_one({"employee_id": emp_id, "date": today})
    pub = employee_public(emp)
    pub["today_status"] = att.get("status") if att else None
    return pub

@api_router.delete("/admin/employees/{emp_id}")
async def delete_employee(emp_id: str, user: dict = Depends(get_current_user)):
    # Pemilik ATAU super admin dapat menghapus pegawai.
    await _get_employee_or_403(emp_id, user)
    await db.employees.delete_one({"id": emp_id})
    await db.attendance.delete_many({"employee_id": emp_id})
    return {"ok": True}

@api_router.post("/admin/attendance")
async def mark_attendance(payload: AttendanceMark, user: dict = Depends(get_current_user)):
    if payload.status not in ATTENDANCE_STATUSES:
        raise HTTPException(status_code=400, detail="Status tidak valid")
    if not _valid_date(payload.date):
        raise HTTPException(status_code=400, detail="Tanggal tidak valid")
    await _get_employee_or_403(payload.employee_id, user, require_owner=True)
    now = datetime.now(timezone.utc).isoformat()
    await db.attendance.update_one(
        {"employee_id": payload.employee_id, "date": payload.date},
        {"$set": {"status": payload.status, "updated_at": now},
         "$setOnInsert": {"id": str(uuid.uuid4()), "employee_id": payload.employee_id, "date": payload.date, "owner_id": user["_id"], "created_at": now}},
        upsert=True,
    )
    return {"ok": True, "employee_id": payload.employee_id, "date": payload.date, "status": payload.status}

@api_router.get("/admin/attendance")
async def list_attendance(employee_id: str, user: dict = Depends(get_current_user)):
    await _get_employee_or_403(employee_id, user)
    docs = await db.attendance.find({"employee_id": employee_id}).sort("date", -1).to_list(400)
    return [{"id": d.get("id"), "date": d.get("date"), "status": d.get("status"), "updated_at": d.get("updated_at")} for d in docs]

# ---------------- Kategori (dinamis) ----------------
DEFAULT_CATEGORY_TILES = {
    "Sepeda Gunung": ("MTB", "https://images.unsplash.com/photo-1594942939850-d8da299577f3?crop=entropy&cs=srgb&fm=jpg&q=85&w=800"),
    "Sepeda Listrik": ("E-Bike", "https://images.unsplash.com/photo-1620802051782-725fa33db067?crop=entropy&cs=srgb&fm=jpg&q=85&w=800"),
    "BMX": ("Freestyle", "https://images.unsplash.com/photo-1628549575837-614973afa6e7?crop=entropy&cs=srgb&fm=jpg&q=85&w=800"),
    "Road Bike": ("Speed", "https://images.unsplash.com/photo-1532298229144-0ec0c57515c7?crop=entropy&cs=srgb&fm=jpg&q=85&w=800"),
    "Sepeda Lipat": ("Compact", "https://images.pexels.com/photos/6558832/pexels-photo-6558832.jpeg?auto=compress&cs=tinysrgb&dpr=2&w=800"),
    "Sepeda Anak": ("Kids", "https://images.unsplash.com/photo-1595182747080-3b43712dd27d?crop=entropy&cs=srgb&fm=jpg&q=85&w=800"),
}

class CategoryCreate(BaseModel):
    name: str
    tag: Optional[str] = ""
    image_url: Optional[str] = ""
    show_on_home: bool = True

class CategoryUpdate(BaseModel):
    name: Optional[str] = None
    tag: Optional[str] = None
    image_url: Optional[str] = None
    show_on_home: Optional[bool] = None
    archived: Optional[bool] = None

class CategoryReorder(BaseModel):
    order: List[str]

def category_public(d: dict, product_count: int = None) -> dict:
    out = {
        "id": d.get("id"),
        "name": d.get("name", ""),
        "tag": d.get("tag", "") or "",
        "image_url": d.get("image_url", "") or "",
        "show_on_home": bool(d.get("show_on_home", False)),
        "archived": bool(d.get("archived", False)),
        "sort_order": int(d.get("sort_order", 0) or 0),
    }
    if product_count is not None:
        out["product_count"] = product_count
    return out

async def active_category_names() -> List[str]:
    docs = await db.categories.find({"archived": {"$ne": True}}).sort("sort_order", 1).to_list(200)
    return [d["name"] for d in docs]

async def seed_categories():
    if await db.categories.count_documents({}) > 0:
        return
    now = datetime.now(timezone.utc).isoformat()
    docs = []
    for i, name in enumerate(CATEGORIES):
        tag, img = DEFAULT_CATEGORY_TILES.get(name, ("", ""))
        docs.append({"id": str(uuid.uuid4()), "name": name, "tag": tag, "image_url": img, "show_on_home": name in DEFAULT_CATEGORY_TILES, "archived": False, "sort_order": i, "created_at": now})
    await db.categories.insert_many(docs)
    logger.info("Categories seeded")

@api_router.get("/admin/categories")
async def admin_list_categories(user: dict = Depends(forbid_cashier)):
    docs = await db.categories.find({}).sort("sort_order", 1).to_list(200)
    counts = {c["_id"]: c["n"] async for c in db.products.aggregate([{"$group": {"_id": "$category", "n": {"$sum": 1}}}])}
    return [category_public(d, counts.get(d.get("name"), 0)) for d in docs]

@api_router.post("/admin/categories")
async def create_category(payload: CategoryCreate, user: dict = Depends(require_super_admin)):
    name = payload.name.strip()
    if not name:
        raise HTTPException(status_code=400, detail="Nama kategori wajib diisi")
    if await db.categories.find_one({"name": {"$regex": f"^{re.escape(name)}$", "$options": "i"}}):
        raise HTTPException(status_code=400, detail="Kategori dengan nama tersebut sudah ada")
    last = await db.categories.find({}).sort("sort_order", -1).to_list(1)
    doc = {
        "id": str(uuid.uuid4()),
        "name": name,
        "tag": (payload.tag or "").strip(),
        "image_url": (payload.image_url or "").strip(),
        "show_on_home": payload.show_on_home,
        "archived": False,
        "sort_order": (int(last[0].get("sort_order", 0)) + 1) if last else 0,
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.categories.insert_one(doc)
    await log_activity(user, "create", "category", f"Tambah kategori {name}", doc["id"])
    return category_public(doc, 0)

@api_router.put("/admin/categories/reorder")
async def reorder_categories(payload: CategoryReorder, user: dict = Depends(require_super_admin)):
    for idx, cid in enumerate(payload.order):
        await db.categories.update_one({"id": cid}, {"$set": {"sort_order": idx}})
    return {"ok": True}

@api_router.put("/admin/categories/{category_id}")
async def update_category(category_id: str, payload: CategoryUpdate, user: dict = Depends(require_super_admin)):
    doc = await db.categories.find_one({"id": category_id})
    if not doc:
        raise HTTPException(status_code=404, detail="Kategori tidak ditemukan")
    updates = {}
    if payload.name is not None:
        name = payload.name.strip()
        if not name:
            raise HTTPException(status_code=400, detail="Nama kategori wajib diisi")
        dup = await db.categories.find_one({"name": {"$regex": f"^{re.escape(name)}$", "$options": "i"}, "id": {"$ne": category_id}})
        if dup:
            raise HTTPException(status_code=400, detail="Kategori dengan nama tersebut sudah ada")
        updates["name"] = name
    for field in ("tag", "image_url"):
        val = getattr(payload, field)
        if val is not None:
            updates[field] = val.strip()
    for field in ("show_on_home", "archived"):
        val = getattr(payload, field)
        if val is not None:
            updates[field] = bool(val)
    if updates:
        await db.categories.update_one({"id": category_id}, {"$set": updates})
        if "name" in updates and updates["name"] != doc.get("name"):
            await db.products.update_many({"category": doc.get("name")}, {"$set": {"category": updates["name"]}})
        if "archived" in updates:
            await log_activity(user, "update", "category", f"{'Arsipkan' if updates['archived'] else 'Aktifkan'} kategori {updates.get('name', doc.get('name'))}", category_id)
    doc = await db.categories.find_one({"id": category_id})
    count = await db.products.count_documents({"category": doc.get("name")})
    return category_public(doc, count)

@api_router.delete("/admin/categories/{category_id}")
async def delete_category(category_id: str, user: dict = Depends(require_super_admin)):
    doc = await db.categories.find_one({"id": category_id})
    if not doc:
        raise HTTPException(status_code=404, detail="Kategori tidak ditemukan")
    count = await db.products.count_documents({"category": doc.get("name")})
    if count > 0:
        raise HTTPException(status_code=400, detail=f"Kategori masih dipakai {count} produk. Arsipkan saja atau pindahkan produknya dulu.")
    await db.categories.delete_one({"id": category_id})
    await log_activity(user, "delete", "category", f"Hapus kategori {doc.get('name')}", category_id)
    return {"ok": True}

# ---------------- Public routes ----------------
@api_router.get("/config")
async def get_config():
    return {"whatsapp_number": WHATSAPP_NUMBER, "categories": await active_category_names()}

@api_router.get("/categories")
async def list_categories_public():
    docs = await db.categories.find({"archived": {"$ne": True}}).sort("sort_order", 1).to_list(200)
    return [category_public(d) for d in docs]

@api_router.get("/products")
async def list_products(category: Optional[str] = None, sort: Optional[str] = None):
    q = {"category": {"$in": await active_category_names()}}
    if category and category != "Semua":
        q["category"] = category
    if sort == "price_asc":
        sort_spec = [("price", 1)]
    elif sort == "price_desc":
        sort_spec = [("price", -1)]
    else:
        sort_spec = [("featured", -1), ("sort_order", 1), ("created_at", -1)]
    docs = await db.products.find(q).sort(sort_spec).to_list(1000)
    return [product_public(d) for d in docs]

@api_router.get("/products/{product_id}")
async def get_product_public(product_id: str):
    try:
        doc = await db.products.find_one({"_id": ObjectId(product_id)})
    except Exception:
        doc = None
    if not doc:
        raise HTTPException(status_code=404, detail="Produk tidak ditemukan")
    return product_public(doc)


@api_router.get("/share/produk/{product_id}")
async def share_product_page(product_id: str, request: Request):
    try:
        doc = await db.products.find_one({"_id": ObjectId(product_id)})
    except Exception:
        doc = None
    if not doc:
        raise HTTPException(status_code=404, detail="Produk tidak ditemukan")
    site = (os.environ.get("FRONTEND_URL") or str(request.base_url)).rstrip("/")
    page_url = f"{site}/produk/{product_id}"
    img = doc.get("image_url") or ""
    if img.startswith("/"):
        img = f"{site}{img}"
    name = html.escape(doc.get("name", "Produk"))
    category = html.escape(doc.get("category", ""))
    desc = html.escape((doc.get("description") or f"{doc.get('name', '')} — {category} tersedia di SK Bike Store Ketapang. Chat admin via WhatsApp untuk info harga & stok.")[:200])
    title = f"{name} | SK Bike Store"
    page = f"""<!doctype html><html lang="id"><head><meta charset="utf-8">
<title>{title}</title>
<meta name="description" content="{desc}">
<link rel="canonical" href="{page_url}">
<meta property="og:type" content="product">
<meta property="og:site_name" content="SK Bike Store">
<meta property="og:locale" content="id_ID">
<meta property="og:title" content="{title}">
<meta property="og:description" content="{desc}">
<meta property="og:url" content="{page_url}">
<meta property="og:image" content="{html.escape(img)}">
<meta property="og:image:alt" content="{name}">
<meta name="twitter:card" content="summary_large_image">
<meta name="twitter:title" content="{title}">
<meta name="twitter:description" content="{desc}">
<meta name="twitter:image" content="{html.escape(img)}">
<meta http-equiv="refresh" content="0;url={page_url}">
<script>window.location.replace({json.dumps(page_url)});</script>
</head><body style="background:#0A0D14;color:#fff;font-family:sans-serif;padding:24px">
<p>Membuka <a href="{page_url}" style="color:#FF2E2E">{name}</a>…</p></body></html>"""
    return StarletteResponse(content=page, media_type="text/html; charset=utf-8")


@api_router.get("/products/image/{path:path}")
async def serve_image(path: str):
    try:
        data, content_type = get_object(path)
        return StarletteResponse(content=data, media_type=content_type)
    except Exception as e:
        logger.error(f"Image serve failed: {e}")
        raise HTTPException(status_code=404, detail="Gambar tidak ditemukan")

# ---------------- Admin routes ----------------
class ReorderPayload(BaseModel):
    order: List[str]

class FeaturedPayload(BaseModel):
    featured: bool

@api_router.get("/admin/products")
async def admin_list_products(user: dict = Depends(get_current_user)):
    docs = await db.products.find({}).sort([("featured", -1), ("sort_order", 1), ("created_at", -1)]).to_list(1000)
    return [product_admin(d, user) for d in docs]

@api_router.put("/admin/products/reorder")
async def reorder_products(payload: ReorderPayload, user: dict = Depends(get_current_user)):
    for idx, pid in enumerate(payload.order):
        try:
            await db.products.update_one({"_id": ObjectId(pid)}, {"$set": {"sort_order": idx}})
        except Exception:
            continue
    return {"ok": True}

@api_router.put("/admin/products/{product_id}/featured")
async def set_product_featured(product_id: str, payload: FeaturedPayload, user: dict = Depends(get_current_user)):
    res = await db.products.update_one({"_id": ObjectId(product_id)}, {"$set": {"featured": bool(payload.featured)}})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Produk tidak ditemukan")
    return {"ok": True, "featured": bool(payload.featured)}

# ---------------- Import / Export produk ----------------
EXPORT_SPEC_KEYS = ["frame", "transmisi", "rem", "ukuran_roda", "baterai_motor"]
EXPORT_COLUMNS = (
    ["name", "code", "category", "description", "price", "cost_price", "status", "image_url"]
    + [f"spec_{k}" for k in EXPORT_SPEC_KEYS]
    + ["variants"]
)

def _variants_to_str(variants) -> str:
    return "; ".join(f"{v.get('color', '')}:{int(v.get('stock', 0) or 0)}" for v in (variants or []) if v.get("color"))

def _parse_variants_str(s) -> list:
    out = []
    for part in str(s or "").split(";"):
        part = part.strip()
        if not part:
            continue
        if ":" in part:
            color, _, stock = part.rpartition(":")
            color = color.strip()
            try:
                stock = int(float(str(stock).strip()))
            except (TypeError, ValueError):
                stock = 0
        else:
            color, stock = part, 0
        if color:
            out.append({"color": color, "stock": max(0, stock)})
    return out

def product_export_row(doc: dict, include_cost: bool = True) -> dict:
    specs = doc.get("specs") or {}
    row = {
        "name": doc.get("name", ""),
        "code": doc.get("code", ""),
        "category": doc.get("category", ""),
        "description": doc.get("description", ""),
        "price": doc.get("price", 0),
        "cost_price": (doc.get("cost_price", 0) if include_cost else ""),
        "status": doc.get("status", "Tersedia"),
        "image_url": doc.get("image_url", ""),
    }
    for k in EXPORT_SPEC_KEYS:
        row[f"spec_{k}"] = specs.get(k, "")
    row["variants"] = _variants_to_str(doc.get("variants"))
    return row

def product_json_item(doc: dict, include_cost: bool = True) -> dict:
    item = {
        "name": doc.get("name", ""),
        "code": doc.get("code", ""),
        "category": doc.get("category", ""),
        "description": doc.get("description", ""),
        "price": doc.get("price", 0),
        "status": doc.get("status", "Tersedia"),
        "image_url": doc.get("image_url", ""),
        "specs": doc.get("specs", {}) or {},
        "variants": doc.get("variants", []) or [],
    }
    if include_cost:
        item["cost_price"] = doc.get("cost_price", 0)
    return item

def _to_num(v) -> float:
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0

def normalize_import_item(item: dict, allowed: List[str]) -> Optional[dict]:
    name = str(item.get("name") or "").strip()
    if not name:
        return None
    category = str(item.get("category") or "").strip()
    if category not in allowed:
        category = allowed[0] if allowed else CATEGORIES[0]
    specs = {}
    src_specs = item.get("specs") if isinstance(item.get("specs"), dict) else {}
    for k in EXPORT_SPEC_KEYS:
        val = src_specs.get(k) if src_specs else item.get(f"spec_{k}")
        specs[k] = str(val or "")
    if isinstance(item.get("variants"), list):
        variants = normalize_variants(item["variants"])
    else:
        variants = _parse_variants_str(item.get("variants"))
    if not variants:
        stock = int(_to_num(item.get("stock")))
        variants = [{"color": "Standar", "stock": max(0, stock)}]
    return {
        "name": name,
        "code": str(item.get("code") or "").strip(),
        "category": category,
        "description": str(item.get("description") or ""),
        "price": _to_num(item.get("price")),
        "cost_price": _to_num(item.get("cost_price")),
        "status": str(item.get("status") or "Tersedia"),
        "image_url": str(item.get("image_url") or ""),
        "specs": specs,
        "variants": variants,
        "stock": sum(v["stock"] for v in variants),
    }

@api_router.get("/admin/products/export")
async def export_products(format: str = "json", user: dict = Depends(forbid_cashier)):
    docs = await db.products.find({}).sort("created_at", -1).to_list(2000)
    show_cost = is_super_admin(user)
    ts = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M")
    if format == "excel":
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Produk"
        ws.append(EXPORT_COLUMNS)
        for d in docs:
            row = product_export_row(d, include_cost=show_cost)
            ws.append([row.get(c, "") for c in EXPORT_COLUMNS])
        for i, col in enumerate(EXPORT_COLUMNS, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(12, min(40, len(col) + 6))
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StarletteResponse(
            content=buf.read(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="produk-skbike-{ts}.xlsx"'},
        )
    payload = json.dumps(
        {"exported_at": datetime.now(timezone.utc).isoformat(), "count": len(docs), "products": [product_json_item(d, include_cost=show_cost) for d in docs]},
        ensure_ascii=False, indent=2,
    )
    return StarletteResponse(
        content=payload.encode("utf-8"),
        media_type="application/json",
        headers={"Content-Disposition": f'attachment; filename="produk-skbike-{ts}.json"'},
    )

@api_router.post("/admin/products/import")
async def import_products(file: UploadFile = File(...), mode: str = Query("merge"), user: dict = Depends(forbid_cashier)):
    if mode not in ("merge", "replace"):
        raise HTTPException(status_code=400, detail="Mode import tidak valid")
    raw = await file.read()
    fname = (file.filename or "").lower()
    items = []
    if fname.endswith(".json"):
        try:
            parsed = json.loads(raw.decode("utf-8"))
        except Exception:
            raise HTTPException(status_code=400, detail="File JSON tidak valid")
        if isinstance(parsed, dict) and "products" in parsed:
            parsed = parsed["products"]
        if not isinstance(parsed, list):
            raise HTTPException(status_code=400, detail="Format JSON harus berupa array produk")
        items = parsed
    elif fname.endswith(".xlsx") or fname.endswith(".xls"):
        import openpyxl
        try:
            wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception:
            raise HTTPException(status_code=400, detail="File Excel tidak dapat dibaca")
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise HTTPException(status_code=400, detail="File Excel kosong")
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        for r in rows[1:]:
            if r is None or all(c is None for c in r):
                continue
            items.append({headers[i]: r[i] for i in range(len(headers)) if i < len(r)})
    elif fname.endswith(".csv"):
        text = raw.decode("utf-8-sig")
        items = list(csv.DictReader(io.StringIO(text)))
    else:
        raise HTTPException(status_code=400, detail="Format file tidak didukung (gunakan .json, .xlsx, atau .csv)")

    allowed = await active_category_names()
    docs = [d for it in items if isinstance(it, dict) for d in [normalize_import_item(it, allowed)] if d]
    if not docs:
        raise HTTPException(status_code=400, detail="Tidak ada produk valid dalam file")

    now = datetime.now(timezone.utc).isoformat()
    created = updated = 0
    if mode == "replace":
        for d in docs:
            d["created_at"] = now
            d["updated_at"] = now
        await db.products.delete_many({})
        await db.products.insert_many(docs)
        created = len(docs)
    else:
        for d in docs:
            existing = None
            if d["code"]:
                existing = await db.products.find_one({"code": d["code"]})
            if existing is None:
                existing = await db.products.find_one({"name": d["name"]})
            if existing:
                d["updated_at"] = now
                await db.products.update_one({"_id": existing["_id"]}, {"$set": d})
                updated += 1
            else:
                d["created_at"] = now
                d["updated_at"] = now
                await db.products.insert_one(d)
                created += 1
    total = await db.products.count_documents({})
    return {"mode": mode, "created": created, "updated": updated, "total": total, "parsed": len(docs)}


# ---------------- Unified Backup & Restore (Super Admin only) ----------------
# Satu file berisi semua data: Penjualan (History), Servis, Pegawai, Absensi, Log Aktivitas.
BACKUP_DATASETS = [
    ("sales", "Penjualan"),
    ("services", "Servis"),
    ("purchase_orders", "Pembelian"),
    ("suppliers", "Supplier"),
    ("categories", "Kategori"),
    ("employees", "Pegawai"),
    ("attendance", "Absensi"),
    ("price_history", "RiwayatHarga"),
    ("pos_transactions", "TransaksiKasir"),
    ("activity_logs", "LogAktivitas"),
]

def _strip_id(d: dict) -> dict:
    return {k: v for k, v in d.items() if k != "_id"}

def _excel_cell(v):
    if v is None:
        return ""
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False)
    if isinstance(v, bool):
        return "true" if v else "false"
    return v

def _excel_uncell(v):
    if isinstance(v, str):
        t = v.strip()
        if t and t[0] in "[{":
            try:
                return json.loads(t)
            except Exception:
                return v
        if t == "true":
            return True
        if t == "false":
            return False
    return v

@api_router.get("/admin/backup/export")
async def backup_export(format: str = "json", user: dict = Depends(require_super_admin)):
    ts = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M")
    data = {}
    for key, _ in BACKUP_DATASETS:
        docs = await db[key].find({}).to_list(100000)
        data[key] = [_strip_id(d) for d in docs]
    # Produk (lossless, termasuk sizes/varian) — disimpan mentah tanpa _id.
    prod_docs = await db.products.find({}).to_list(100000)
    data["products"] = [_strip_id(d) for d in prod_docs]

    if format == "excel":
        import openpyxl
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        sheets = BACKUP_DATASETS + [("products", "Produk")]
        for key, sheet_name in sheets:
            ws = wb.create_sheet(title=sheet_name[:31])
            rows = data[key]
            cols = []
            for r in rows:
                for c in r.keys():
                    if c not in cols:
                        cols.append(c)
            if not cols:
                cols = ["id"]
            ws.append(cols)
            for r in rows:
                ws.append([_excel_cell(r.get(c)) for c in cols])
            for i, col in enumerate(cols, start=1):
                ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(12, min(40, len(str(col)) + 6))
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StarletteResponse(
            content=buf.read(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="backup-skbike-{ts}.xlsx"'},
        )

    payload = json.dumps({
        "type": "skbike-backup",
        "version": 2,
        "exported_at": datetime.now(timezone.utc).isoformat(),
        "counts": {k: len(v) for k, v in data.items()},
        "data": data,
    }, ensure_ascii=False, indent=2, default=str)
    return StarletteResponse(
        content=payload.encode("utf-8"),
        media_type="application/json",
        headers={"Content-Disposition": f'attachment; filename="backup-skbike-{ts}.json"'},
    )

@api_router.post("/admin/backup/import")
async def backup_import(file: UploadFile = File(...), mode: str = Query("merge"), user: dict = Depends(require_super_admin)):
    if mode not in ("merge", "replace"):
        raise HTTPException(status_code=400, detail="Mode import tidak valid")
    raw = await file.read()
    fname = (file.filename or "").lower()
    data = {}
    all_keys = [k for k, _ in BACKUP_DATASETS] + ["products"]

    if fname.endswith(".json"):
        try:
            parsed = json.loads(raw.decode("utf-8"))
        except Exception:
            raise HTTPException(status_code=400, detail="File JSON tidak valid")
        blob = parsed.get("data") if isinstance(parsed, dict) and "data" in parsed else parsed
        if not isinstance(blob, dict):
            raise HTTPException(status_code=400, detail="Struktur backup JSON tidak dikenali")
        for key in all_keys:
            rows = blob.get(key) or []
            data[key] = [r for r in rows if isinstance(r, dict)]
    elif fname.endswith(".xlsx") or fname.endswith(".xls"):
        import openpyxl
        try:
            wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception:
            raise HTTPException(status_code=400, detail="File Excel tidak dapat dibaca")
        name_to_key = {sheet_name: key for key, sheet_name in (BACKUP_DATASETS + [("products", "Produk")])}
        for key in all_keys:
            data[key] = []
        for ws in wb.worksheets:
            key = name_to_key.get(ws.title)
            if not key:
                continue
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            headers = [str(h).strip() if h is not None else "" for h in rows[0]]
            for r in rows[1:]:
                if r is None or all(c is None for c in r):
                    continue
                item = {headers[i]: _excel_uncell(r[i]) for i in range(len(headers)) if i < len(r) and headers[i]}
                data[key].append(item)
    else:
        raise HTTPException(status_code=400, detail="Format file tidak didukung (gunakan .json atau .xlsx)")

    result = {}
    now = datetime.now(timezone.utc).isoformat()
    for key, _ in BACKUP_DATASETS:
        rows = data.get(key) or []
        created = updated = 0
        clean = [_strip_id(r) for r in rows if isinstance(r, dict)]
        if mode == "replace":
            await db[key].delete_many({})
            if clean:
                await db[key].insert_many([dict(r) for r in clean])
                created = len(clean)
        else:
            for r in clean:
                rid = r.get("id")
                if rid:
                    existing = await db[key].find_one({"id": rid})
                    if existing:
                        await db[key].update_one({"id": rid}, {"$set": r})
                        updated += 1
                        continue
                elif await db[key].find_one(dict(r)):
                    # koleksi tanpa "id" (mis. price_history): lewati duplikat identik
                    updated += 1
                    continue
                r.setdefault("created_at", now)
                await db[key].insert_one(dict(r))
                created += 1
        total = await db[key].count_documents({})
        result[key] = {"created": created, "updated": updated, "total": total, "parsed": len(clean)}

    # Produk: cocokkan berdasarkan kode lalu nama (menjaga sizes/varian utuh).
    prod_rows = data.get("products") or []
    p_created = p_updated = 0
    p_clean = [_strip_id(r) for r in prod_rows if isinstance(r, dict) and (r.get("name"))]
    if mode == "replace":
        await db.products.delete_many({})
        if p_clean:
            for r in p_clean:
                r.setdefault("created_at", now)
                r["updated_at"] = now
            await db.products.insert_many([dict(r) for r in p_clean])
            p_created = len(p_clean)
    else:
        for r in p_clean:
            existing = None
            if r.get("code"):
                existing = await db.products.find_one({"code": r["code"]})
            if existing is None:
                existing = await db.products.find_one({"name": r["name"]})
            r["updated_at"] = now
            if existing:
                await db.products.update_one({"_id": existing["_id"]}, {"$set": r})
                p_updated += 1
            else:
                r.setdefault("created_at", now)
                await db.products.insert_one(dict(r))
                p_created += 1
    result["products"] = {"created": p_created, "updated": p_updated, "total": await db.products.count_documents({}), "parsed": len(p_clean)}

    await log_activity(user, "update", "account", f"Restore backup ({mode})")
    return {"mode": mode, "datasets": result}


@api_router.get("/admin/stats")
async def admin_stats(user: dict = Depends(get_current_user)):
    docs = await db.products.find({}).to_list(1000)
    total = len(docs)
    cats = len(set(d.get("category") for d in docs))
    in_stock = sum(1 for d in docs if d.get("stock", 0) > 0)
    inventory_value = sum(d.get("price", 0) * d.get("stock", 0) for d in docs)
    return {"total_products": total, "total_categories": cats, "in_stock": in_stock, "inventory_value": inventory_value}

LOW_STOCK_THRESHOLD = 3

@api_router.get("/admin/low-stock")
async def low_stock(user: dict = Depends(get_current_user)):
    docs = await db.products.find({}).sort("stock", 1).to_list(1000)
    out_items, low_items = [], []
    for d in docs:
        stock = int(d.get("stock", 0) or 0)
        item = {
            "id": str(d["_id"]),
            "name": d.get("name"),
            "code": d.get("code", ""),
            "category": d.get("category"),
            "stock": stock,
            "image_url": d.get("image_url", ""),
        }
        if stock <= 0:
            out_items.append(item)
        elif stock <= LOW_STOCK_THRESHOLD:
            low_items.append(item)
    return {
        "threshold": LOW_STOCK_THRESHOLD,
        "out_of_stock": out_items,
        "low_stock": low_items,
        "count": len(out_items) + len(low_items),
    }

# ---------------- Supplier & Purchase Order (Restock) ----------------
class SupplierCreate(BaseModel):
    nama: str
    telepon: Optional[str] = None
    alamat: Optional[str] = None
    catatan: Optional[str] = None

class SupplierUpdate(BaseModel):
    nama: Optional[str] = None
    telepon: Optional[str] = None
    alamat: Optional[str] = None
    catatan: Optional[str] = None

class POItem(BaseModel):
    product_id: Optional[str] = None
    name: str
    code: Optional[str] = ""
    color: Optional[str] = None
    size: Optional[str] = None
    qty: int = 1
    unit_cost: float = 0

class PurchaseOrderCreate(BaseModel):
    supplier_id: Optional[str] = None
    supplier_name: Optional[str] = None
    tanggal: Optional[str] = None
    catatan: Optional[str] = None
    update_cost: bool = True
    status: Optional[str] = "diterima"
    items: List[POItem] = []

def supplier_public(d: dict) -> dict:
    return {
        "id": d.get("id"),
        "nama": d.get("nama"),
        "telepon": d.get("telepon"),
        "alamat": d.get("alamat"),
        "catatan": d.get("catatan"),
        "created_at": d.get("created_at"),
    }

def po_public(d: dict) -> dict:
    items = []
    default_recv_full = d.get("status") == "diterima"
    for it in d.get("items", []):
        it = dict(it)
        if "received_qty" not in it:
            it["received_qty"] = int(it.get("qty", 0)) if default_recv_full else 0
        items.append(it)
    return {
        "id": d.get("id"),
        "supplier_id": d.get("supplier_id"),
        "supplier_name": d.get("supplier_name"),
        "tanggal": d.get("tanggal"),
        "catatan": d.get("catatan"),
        "items": items,
        "total": d.get("total", 0),
        "status": d.get("status", "diterima"),
        "paid": bool(d.get("paid", False)),
        "paid_at": d.get("paid_at"),
        "update_cost": d.get("update_cost", True),
        "received_at": d.get("received_at"),
        "created_by": d.get("created_by"),
        "created_at": d.get("created_at"),
    }

def po_status_from_items(items) -> str:
    total = sum(int(i.get("qty", 0)) for i in items)
    recv = sum(int(i.get("received_qty", 0)) for i in items)
    if recv <= 0:
        return "dipesan"
    if recv >= total:
        return "diterima"
    return "sebagian"

async def add_stock(product_id: str, qty: int, color: Optional[str] = None, size: Optional[str] = None, unit_cost: Optional[float] = None, update_cost: bool = False):
    prod = await db.products.find_one({"_id": ObjectId(product_id)})
    if not prod:
        return
    sizes = prod.get("sizes") or []
    if sizes:
        size_obj = next((s for s in sizes if s.get("label") == size), None) if size else None
        if size_obj is None:
            size_obj = sizes[0]
        cols = size_obj.get("colors") or []
        if cols:
            target = next((c for c in cols if c.get("color") == color), None) if color else None
            if target is None:
                target = cols[0]
            target["stock"] = int(target.get("stock", 0)) + qty
        if update_cost and unit_cost:
            size_obj["cost_price"] = float(unit_cost)
        agg = aggregate_variants_from_sizes(sizes)
        total = sum(v["stock"] for v in agg)
        upd = {"sizes": sizes, "variants": agg, "stock": total}
        if update_cost and sizes:
            upd["cost_price"] = min(s.get("cost_price", 0) for s in sizes)
            upd["price"] = min(s.get("price", 0) for s in sizes)
        await db.products.update_one({"_id": ObjectId(product_id)}, {"$set": upd})
        return
    variants = prod.get("variants") or []
    if variants:
        target = next((v for v in variants if v.get("color") == color), None) if color else None
        if target is None:
            target = variants[0]
        target["stock"] = int(target.get("stock", 0)) + qty
        total = sum(int(v.get("stock", 0)) for v in variants)
        upd = {"variants": variants, "stock": total}
        if update_cost and unit_cost:
            upd["cost_price"] = float(unit_cost)
        await db.products.update_one({"_id": ObjectId(product_id)}, {"$set": upd})
    else:
        await db.products.update_one({"_id": ObjectId(product_id)}, {"$inc": {"stock": qty}})
        if update_cost and unit_cost:
            await db.products.update_one({"_id": ObjectId(product_id)}, {"$set": {"cost_price": float(unit_cost)}})

@api_router.get("/admin/suppliers")
async def list_suppliers(user: dict = Depends(forbid_cashier)):
    docs = await db.suppliers.find({}).sort("nama", 1).to_list(1000)
    return [supplier_public(d) for d in docs]

@api_router.post("/admin/suppliers")
async def create_supplier(payload: SupplierCreate, user: dict = Depends(forbid_cashier)):
    if not payload.nama.strip():
        raise HTTPException(status_code=400, detail="Nama supplier wajib diisi")
    doc = {
        "id": str(uuid.uuid4()),
        "nama": payload.nama.strip(),
        "telepon": (payload.telepon or "").strip() or None,
        "alamat": (payload.alamat or "").strip() or None,
        "catatan": (payload.catatan or "").strip() or None,
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.suppliers.insert_one(doc)
    await log_activity(user, "create", "supplier", f"Tambah supplier {doc['nama']}", doc["id"])
    return supplier_public(doc)

@api_router.put("/admin/suppliers/{supplier_id}")
async def update_supplier(supplier_id: str, payload: SupplierUpdate, user: dict = Depends(forbid_cashier)):
    doc = await db.suppliers.find_one({"id": supplier_id})
    if not doc:
        raise HTTPException(status_code=404, detail="Supplier tidak ditemukan")
    updates = {}
    if payload.nama is not None:
        if not payload.nama.strip():
            raise HTTPException(status_code=400, detail="Nama supplier wajib diisi")
        updates["nama"] = payload.nama.strip()
    for field in ("telepon", "alamat", "catatan"):
        val = getattr(payload, field)
        if val is not None:
            updates[field] = val.strip() or None
    if updates:
        await db.suppliers.update_one({"id": supplier_id}, {"$set": updates})
    doc = await db.suppliers.find_one({"id": supplier_id})
    return supplier_public(doc)

@api_router.delete("/admin/suppliers/{supplier_id}")
async def delete_supplier(supplier_id: str, user: dict = Depends(forbid_cashier)):
    doc = await db.suppliers.find_one({"id": supplier_id})
    if not doc:
        raise HTTPException(status_code=404, detail="Supplier tidak ditemukan")
    await db.suppliers.delete_one({"id": supplier_id})
    await log_activity(user, "delete", "supplier", f"Hapus supplier {doc.get('nama')}", supplier_id)
    return {"ok": True}

@api_router.get("/admin/suppliers/{supplier_id}/export")
async def export_supplier(supplier_id: str, user: dict = Depends(forbid_cashier)):
    sup = await db.suppliers.find_one({"id": supplier_id})
    if not sup:
        raise HTTPException(status_code=404, detail="Supplier tidak ditemukan")
    nama = (sup.get("nama") or "").strip()
    products = await db.products.find({"supplier": nama}).sort("name", 1).to_list(5000)
    orders = await db.purchase_orders.find({"supplier_id": supplier_id}).sort("created_at", -1).to_list(5000)
    show_cost = is_super_admin(user)

    import openpyxl
    wb = openpyxl.Workbook()

    # Sheet 1: Info supplier
    ws_info = wb.active
    ws_info.title = "Info"
    pending = [o for o in orders if o.get("status", "diterima") in ("dipesan", "sebagian")]
    info_rows = [
        ("Nama Supplier", nama),
        ("Telepon", sup.get("telepon") or "-"),
        ("Alamat", sup.get("alamat") or "-"),
        ("Catatan", sup.get("catatan") or "-"),
        ("Jumlah Produk", len(products)),
        ("Jumlah Pembelian", len(orders)),
        ("Pembelian Menunggu Barang", len(pending)),
        ("Total Belanja", sum(float(o.get("total", 0) or 0) for o in orders)),
    ]
    for r in info_rows:
        ws_info.append(list(r))
    ws_info.column_dimensions["A"].width = 28
    ws_info.column_dimensions["B"].width = 45

    # Sheet 2: Produk
    ws_prod = wb.create_sheet("Produk")
    prod_cols = ["Nama", "Kode", "Kategori", "Stok", "Harga Jual", "Status"]
    if show_cost:
        prod_cols.insert(5, "Harga Modal")
    ws_prod.append(prod_cols)
    for p in products:
        row = [p.get("name", ""), p.get("code", ""), p.get("category", ""), int(p.get("stock", 0) or 0), float(p.get("price", 0) or 0)]
        if show_cost:
            row.append(float(p.get("cost_price", 0) or 0))
        row.append(p.get("status", "Tersedia"))
        ws_prod.append(row)
    for i, col in enumerate(prod_cols, start=1):
        ws_prod.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(12, min(40, len(col) + 8))

    # Sheet 3: Pembelian
    ws_po = wb.create_sheet("Pembelian")
    po_cols = ["Tanggal", "Status", "Jumlah Item", "Total", "Catatan", "Dibuat Oleh", "Rincian Item"]
    ws_po.append(po_cols)
    status_label = {"dipesan": "Dipesan", "sebagian": "Sebagian", "diterima": "Diterima"}
    for o in orders:
        detail = "; ".join(
            f"{it.get('name', '')}"
            + (f" ({it.get('size')})" if it.get("size") else "")
            + (f" [{it.get('color')}]" if it.get("color") else "")
            + f" x{int(it.get('qty', 0) or 0)} @ {int(it.get('unit_cost', 0) or 0)}"
            for it in (o.get("items") or [])
        )
        ws_po.append([
            o.get("tanggal", ""),
            status_label.get(o.get("status", "diterima"), o.get("status", "")),
            len(o.get("items") or []),
            float(o.get("total", 0) or 0),
            o.get("catatan") or "",
            o.get("created_by") or "",
            detail,
        ])
    for i, col in enumerate(po_cols, start=1):
        ws_po.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(12, min(60, len(col) + 8))

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    safe = "".join(c for c in nama if c.isalnum() or c in (" ", "-", "_")).strip().replace(" ", "-") or "supplier"
    ts = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M")
    return StarletteResponse(
        content=buf.read(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="supplier-{safe}-{ts}.xlsx"'},
    )

@api_router.get("/admin/purchase-orders")
async def list_purchase_orders(supplier_id: Optional[str] = None, user: dict = Depends(forbid_cashier)):
    q = {}
    if supplier_id:
        q["supplier_id"] = supplier_id
    docs = await db.purchase_orders.find(q).sort("created_at", -1).to_list(2000)
    return [po_public(d) for d in docs]

@api_router.get("/admin/purchase-orders/summary")
async def purchase_orders_summary(user: dict = Depends(forbid_cashier)):
    docs = await db.purchase_orders.find({}).to_list(5000)
    wib = timezone(timedelta(hours=7))
    month_key = datetime.now(wib).strftime("%Y-%m")
    def po_month(d):
        raw = d.get("created_at")
        try:
            dt = datetime.fromisoformat(str(raw))
            if dt.tzinfo is None:
                dt = dt.replace(tzinfo=timezone.utc)
            return dt.astimezone(wib).strftime("%Y-%m")
        except (TypeError, ValueError):
            return ""
    month_docs = [d for d in docs if po_month(d) == month_key]
    pending = [d for d in docs if d.get("status", "diterima") in ("dipesan", "sebagian")]
    return {
        "total_orders": len(docs),
        "total_spend": sum(float(d.get("total", 0) or 0) for d in docs),
        "month_orders": len(month_docs),
        "month_spend": sum(float(d.get("total", 0) or 0) for d in month_docs),
        "supplier_count": await db.suppliers.count_documents({}),
        "pending_orders": len(pending),
        "pending_spend": sum(float(d.get("total", 0) or 0) for d in pending),
    }

@api_router.post("/admin/purchase-orders")
async def create_purchase_order(payload: PurchaseOrderCreate, user: dict = Depends(forbid_cashier)):
    if not payload.items:
        raise HTTPException(status_code=400, detail="Item pembelian kosong")
    status = (payload.status or "diterima").lower()
    if status not in ("dipesan", "diterima"):
        status = "diterima"
    supplier_name = (payload.supplier_name or "").strip() or None
    if payload.supplier_id:
        sup = await db.suppliers.find_one({"id": payload.supplier_id})
        if sup:
            supplier_name = sup.get("nama")
    now_dt = datetime.now(timezone.utc)
    items = []
    total = 0.0
    for it in payload.items:
        qty = max(1, int(it.qty or 1))
        cost = max(0.0, float(it.unit_cost or 0))
        line = qty * cost
        total += line
        received_qty = qty if status == "diterima" else 0
        items.append({
            "id": str(uuid.uuid4()),
            "product_id": it.product_id, "name": it.name, "code": it.code or "",
            "color": it.color, "size": it.size, "qty": qty, "received_qty": received_qty,
            "unit_cost": cost, "line_total": line,
        })
        # Stok hanya bertambah bila barang sudah diterima (sampai di gudang/toko).
        if status == "diterima" and it.product_id:
            try:
                await add_stock(it.product_id, qty, it.color, it.size, cost, payload.update_cost)
            except Exception as e:
                logger.error(f"Restock failed for {it.product_id}: {e}")
    doc = {
        "id": str(uuid.uuid4()),
        "supplier_id": payload.supplier_id,
        "supplier_name": supplier_name,
        "tanggal": (payload.tanggal or "").strip() or id_date_label(now_dt),
        "catatan": (payload.catatan or "").strip() or None,
        "items": items,
        "total": total,
        "status": status,
        "update_cost": bool(payload.update_cost),
        "received_at": now_dt.isoformat() if status == "diterima" else None,
        "created_by": user.get("email", ""),
        "created_at": now_dt.isoformat(),
    }
    await db.purchase_orders.insert_one(dict(doc))
    label = "Pembelian" if status == "diterima" else "Pesanan"
    await log_activity(user, "create", "purchase", f"{label} {len(items)} item • {rupiah_id(total)}" + (f" dari {supplier_name}" if supplier_name else ""), doc["id"])
    return po_public(doc)

class POReceiveItem(BaseModel):
    item_id: str
    qty: int

class POReceive(BaseModel):
    items: Optional[List[POReceiveItem]] = None

@api_router.post("/admin/purchase-orders/{po_id}/receive")
async def receive_purchase_order(po_id: str, payload: Optional[POReceive] = None, user: dict = Depends(forbid_cashier)):
    doc = await db.purchase_orders.find_one({"id": po_id})
    if not doc:
        raise HTTPException(status_code=404, detail="Pembelian tidak ditemukan")
    if doc.get("status") == "diterima":
        raise HTTPException(status_code=400, detail="Pembelian sudah diterima seluruhnya")
    items = doc.get("items", [])
    for it in items:
        if "id" not in it:
            it["id"] = str(uuid.uuid4())
        it.setdefault("received_qty", 0)
    by_id = {it["id"]: it for it in items}
    update_cost = bool(doc.get("update_cost", True))
    if payload and payload.items:
        reqs = [(r.item_id, int(r.qty)) for r in payload.items]
    else:
        reqs = [(it["id"], int(it.get("qty", 0)) - int(it.get("received_qty", 0))) for it in items]
    any_recv = False
    for item_id, want in reqs:
        it = by_id.get(item_id)
        if not it:
            continue
        remaining = int(it.get("qty", 0)) - int(it.get("received_qty", 0))
        recv = max(0, min(int(want), remaining))
        if recv <= 0:
            continue
        any_recv = True
        if it.get("product_id"):
            try:
                await add_stock(it["product_id"], recv, it.get("color"), it.get("size"), float(it.get("unit_cost") or 0), update_cost)
            except Exception as e:
                logger.error(f"Receive restock failed for {it.get('product_id')}: {e}")
        it["received_qty"] = int(it.get("received_qty", 0)) + recv
    if not any_recv:
        raise HTTPException(status_code=400, detail="Tidak ada barang untuk diterima")
    new_status = po_status_from_items(items)
    updates = {"items": items, "status": new_status}
    if new_status == "diterima":
        updates["received_at"] = datetime.now(timezone.utc).isoformat()
    await db.purchase_orders.update_one({"id": po_id}, {"$set": updates})
    await log_activity(user, "update", "purchase", f"Terima barang ({new_status}) • {doc.get('supplier_name') or 'PO'}", po_id)
    doc = await db.purchase_orders.find_one({"id": po_id})
    return po_public(doc)

@api_router.delete("/admin/purchase-orders/{po_id}")
async def delete_purchase_order(po_id: str, user: dict = Depends(forbid_cashier)):
    doc = await db.purchase_orders.find_one({"id": po_id})
    if not doc:
        raise HTTPException(status_code=404, detail="Pembelian tidak ditemukan")
    # Kembalikan stok sebanyak yang sudah benar-benar diterima per item.
    default_full = doc.get("status") == "diterima"
    for it in doc.get("items", []):
        recv = int(it.get("received_qty", it.get("qty", 0) if default_full else 0))
        if it.get("product_id") and recv > 0:
            try:
                await deduct_stock(it["product_id"], recv, it.get("color"), it.get("size"))
            except Exception as e:
                logger.error(f"Reverse restock failed for {it.get('product_id')}: {e}")
    await db.purchase_orders.delete_one({"id": po_id})
    await log_activity(user, "delete", "purchase", f"Hapus pembelian {rupiah_id(doc.get('total', 0))}", po_id)
    return {"ok": True}

class POPaid(BaseModel):
    paid: bool

@api_router.patch("/admin/purchase-orders/{po_id}/paid")
async def set_purchase_order_paid(po_id: str, payload: POPaid, user: dict = Depends(forbid_cashier)):
    doc = await db.purchase_orders.find_one({"id": po_id})
    if not doc:
        raise HTTPException(status_code=404, detail="Pembelian tidak ditemukan")
    paid = bool(payload.paid)
    await db.purchase_orders.update_one({"id": po_id}, {"$set": {"paid": paid, "paid_at": datetime.now(timezone.utc).isoformat() if paid else None}})
    doc = await db.purchase_orders.find_one({"id": po_id})
    return po_public(doc)

# ---------------- Kasir / Sales routes ----------------
@api_router.post("/admin/sales")
async def create_sale(payload: SaleCreate, user: dict = Depends(get_current_user)):
    now = datetime.now(timezone.utc).isoformat()
    doc = payload.dict()
    doc["id"] = str(uuid.uuid4())
    doc["created_at"] = now
    doc["deleted_at"] = None
    doc["cashier"] = user.get("email", "")
    await db.sales.insert_one(doc)
    await log_activity(user, "create", "sale", f"Penjualan {doc.get('nama_barang') or 'barang'}", doc["id"])
    return sale_public(doc)

@api_router.get("/admin/sales")
async def list_sales(user: dict = Depends(get_current_user)):
    q = {"deleted_at": None}
    if not is_super_admin(user):
        q["cashier"] = user["email"]
    docs = await db.sales.find(q).sort("created_at", -1).to_list(1000)
    emails = list({d.get("cashier") for d in docs if d.get("cashier")})
    name_map = {}
    if emails:
        users = await db.users.find({"email": {"$in": emails}}).to_list(200)
        name_map = {u["email"]: u.get("name", "Admin") for u in users}
    result = []
    for d in docs:
        pub = sale_public(d)
        pub["owner_name"] = name_map.get(pub["owner_email"], pub["owner_email"] or "Admin")
        result.append(strip_sale_financials(pub, user))
    return result

def _num(d: dict, k: str) -> float:
    v = d.get(k)
    try:
        return float(v) if v else 0.0
    except (TypeError, ValueError):
        return 0.0

@api_router.get("/admin/sales/summary")
async def sales_summary(user: dict = Depends(get_current_user)):
    q = {"deleted_at": None, "returned": {"$ne": True}}
    if not is_super_admin(user):
        q["cashier"] = user["email"]
    docs = await db.sales.find(q).to_list(5000)
    today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    today_docs = [d for d in docs if str(d.get("created_at", ""))[:10] == today]
    if not can_see_financials(user):
        return {
            "total_transactions": len(docs),
            "total_revenue": None,
            "total_margin": None,
            "today_transactions": len(today_docs),
            "today_revenue": None,
            "today_margin": None,
        }
    return {
        "total_transactions": len(docs),
        "total_revenue": sum(_num(d, "harga_jual") for d in docs),
        "total_margin": sum(_num(d, "margin") for d in docs),
        "today_transactions": len(today_docs),
        "today_revenue": sum(_num(d, "harga_jual") for d in today_docs),
        "today_margin": sum(_num(d, "margin") for d in today_docs),
    }

@api_router.get("/admin/sales/daily")
async def sales_daily(days: int = 14, user: dict = Depends(get_current_user)):
    if not can_see_financials(user):
        raise HTTPException(status_code=403, detail="Akses ditolak: data omset & margin dibatasi")
    days = max(1, min(int(days or 14), 90))
    q = {"deleted_at": None, "returned": {"$ne": True}}
    if not is_super_admin(user):
        q["cashier"] = user["email"]
    docs = await db.sales.find(q).to_list(10000)
    # Group by local (WIB, UTC+7) calendar day.
    wib = timezone(timedelta(hours=7))
    buckets = {}
    for d in docs:
        raw = d.get("created_at")
        if not raw:
            continue
        try:
            dt = datetime.fromisoformat(str(raw))
            if dt.tzinfo is None:
                dt = dt.replace(tzinfo=timezone.utc)
            key = dt.astimezone(wib).strftime("%Y-%m-%d")
        except (TypeError, ValueError):
            continue
        b = buckets.setdefault(key, {"revenue": 0.0, "margin": 0.0, "transactions": 0})
        b["revenue"] += _num(d, "harga_jual")
        b["margin"] += _num(d, "margin")
        b["transactions"] += 1
    today = datetime.now(wib).date()
    out = []
    for i in range(days - 1, -1, -1):
        day = today - timedelta(days=i)
        key = day.strftime("%Y-%m-%d")
        b = buckets.get(key, {"revenue": 0.0, "margin": 0.0, "transactions": 0})
        out.append({
            "date": key,
            "label": day.strftime("%d %b"),
            "revenue": round(b["revenue"], 2),
            "margin": round(b["margin"], 2),
            "transactions": b["transactions"],
        })
    return out

@api_router.get("/admin/dashboard")
async def sales_dashboard(period: str = Query("today", alias="range"), start: str = None, end: str = None, user: dict = Depends(get_current_user)):
    if not can_see_financials(user):
        raise HTTPException(status_code=403, detail="Akses ditolak: data omset & laba dibatasi")
    wib = timezone(timedelta(hours=7))
    today = datetime.now(wib).date()
    # Resolve date window [start_date, end_date] inclusive (WIB calendar days).
    rng = (period or "today").lower()
    if rng == "today":
        start_date = end_date = today
    elif rng == "yesterday":
        start_date = end_date = today - timedelta(days=1)
    elif rng in ("7d", "7days", "week"):
        start_date, end_date = today - timedelta(days=6), today
    elif rng in ("30d", "30days"):
        start_date, end_date = today - timedelta(days=29), today
    elif rng in ("month", "this_month"):
        start_date, end_date = today.replace(day=1), today
    elif rng == "custom":
        try:
            start_date = datetime.strptime(start, "%Y-%m-%d").date() if start else today
            end_date = datetime.strptime(end, "%Y-%m-%d").date() if end else today
        except (TypeError, ValueError):
            raise HTTPException(status_code=400, detail="Tanggal custom tidak valid (format YYYY-MM-DD)")
        if end_date < start_date:
            start_date, end_date = end_date, start_date
    else:
        start_date = end_date = today

    q = {"deleted_at": None, "returned": {"$ne": True}}
    if not is_super_admin(user):
        q["cashier"] = user["email"]
    docs = await db.sales.find(q).to_list(20000)

    def day_key(d):
        raw = d.get("created_at")
        if not raw:
            return None
        try:
            dt = datetime.fromisoformat(str(raw))
            if dt.tzinfo is None:
                dt = dt.replace(tzinfo=timezone.utc)
            return dt.astimezone(wib).date()
        except (TypeError, ValueError):
            return None

    revenue = profit = cost = units = 0.0
    tx_ids = set()
    span_days = (end_date - start_date).days + 1
    buckets = {}
    prod_agg = {}
    pay_agg = {}
    recent_pool = []
    for d in docs:
        dk = day_key(d)
        if dk is None or dk < start_date or dk > end_date:
            continue
        r = _num(d, "harga_jual")
        m = _num(d, "margin")
        c = _num(d, "harga_modal")
        qty = int(d.get("qty") or 0) or 1
        revenue += r
        profit += m
        cost += c
        units += qty
        tx_ids.add(d.get("transaction_id") or d.get("id"))
        # Produk terlaris (agregasi per nama barang)
        pname = (str(d.get("nama_barang") or d.get("kode_barang") or "Tanpa Nama").strip()) or "Tanpa Nama"
        pa = prod_agg.setdefault(pname, {"name": pname, "units": 0, "revenue": 0.0})
        pa["units"] += qty
        pa["revenue"] += r
        # Metode pembayaran (Cash / Transfer / Split / Lainnya)
        method = (str(d.get("metode_pembayaran") or "Lainnya").strip()) or "Lainnya"
        pm = pay_agg.setdefault(method, {"method": method, "count": 0, "revenue": 0.0})
        pm["count"] += 1
        pm["revenue"] += r
        b = buckets.setdefault(dk.strftime("%Y-%m-%d"), {"revenue": 0.0, "profit": 0.0, "cost": 0.0, "units": 0, "transactions": set()})
        b["revenue"] += r
        b["profit"] += m
        b["cost"] += c
        b["units"] += qty
        b["transactions"].add(d.get("transaction_id") or d.get("id"))
        recent_pool.append({
            "id": d.get("id"),
            "nama_barang": d.get("nama_barang") or d.get("kode_barang") or "Tanpa Nama",
            "nama_pembeli": d.get("nama_pembeli") or "",
            "harga_jual": round(r, 2),
            "metode_pembayaran": d.get("metode_pembayaran") or "Lainnya",
            "sudah_diambil": d.get("sudah_diambil") or "Belum",
            "created_at": d.get("created_at"),
            "_dk": dk,
        })

    transactions = len(tx_ids)
    avg_tx = (revenue / transactions) if transactions else 0.0

    top_products = sorted(prod_agg.values(), key=lambda x: x["revenue"], reverse=True)[:10]
    for t in top_products:
        t["revenue"] = round(t["revenue"], 2)
    payment_breakdown = sorted(pay_agg.values(), key=lambda x: x["revenue"], reverse=True)
    for p in payment_breakdown:
        p["revenue"] = round(p["revenue"], 2)

    recent = sorted(recent_pool, key=lambda x: str(x.get("created_at") or ""), reverse=True)[:8]
    for rc in recent:
        rc.pop("_dk", None)

    # Perbandingan dengan periode sebelumnya (jendela sama panjang, tepat sebelum start_date)
    prev_end = start_date - timedelta(days=1)
    prev_start = prev_end - timedelta(days=span_days - 1)
    prev_rev = prev_profit = prev_cost = prev_units = 0.0
    prev_tx = set()
    for d in docs:
        dk = day_key(d)
        if dk is None or dk < prev_start or dk > prev_end:
            continue
        prev_rev += _num(d, "harga_jual")
        prev_profit += _num(d, "margin")
        prev_cost += _num(d, "harga_modal")
        prev_units += int(d.get("qty") or 0) or 1
        prev_tx.add(d.get("transaction_id") or d.get("id"))
    prev_transactions = len(prev_tx)
    prev_avg = (prev_rev / prev_transactions) if prev_transactions else 0.0
    previous = {
        "revenue": round(prev_rev, 2),
        "gross_profit": round(prev_profit, 2),
        "cost": round(prev_cost, 2),
        "transactions": prev_transactions,
        "units_sold": int(prev_units),
        "avg_transaction_value": round(prev_avg, 2),
        "start": prev_start.strftime("%Y-%m-%d"),
        "end": prev_end.strftime("%Y-%m-%d"),
    }

    chart = []
    for i in range(span_days):
        day = start_date + timedelta(days=i)
        key = day.strftime("%Y-%m-%d")
        b = buckets.get(key, {"revenue": 0.0, "profit": 0.0, "cost": 0.0, "units": 0, "transactions": set()})
        tx_count = len(b["transactions"]) if isinstance(b["transactions"], set) else 0
        rev = round(b["revenue"], 2)
        chart.append({
            "date": key,
            "label": day.strftime("%d %b"),
            "revenue": rev,
            "profit": round(b["profit"], 2),
            "cost": round(b["cost"], 2),
            "units": int(b["units"]),
            "transactions": tx_count,
            "avg": round(rev / tx_count, 2) if tx_count else 0,
        })

    return {
        "range": rng,
        "start": start_date.strftime("%Y-%m-%d"),
        "end": end_date.strftime("%Y-%m-%d"),
        "summary": {
            "revenue": round(revenue, 2),
            "gross_profit": round(profit, 2),
            "cost": round(cost, 2),
            "transactions": transactions,
            "units_sold": int(units),
            "avg_transaction_value": round(avg_tx, 2),
        },
        "previous": previous,
        "top_products": top_products,
        "payment_breakdown": payment_breakdown,
        "recent": recent,
        "chart": chart,
    }


@api_router.put("/admin/sales/{sale_id}")
async def update_sale(sale_id: str, payload: SaleCreate, user: dict = Depends(get_current_user)):
    doc = await db.sales.find_one({"id": sale_id, "deleted_at": None})
    if not doc:
        raise HTTPException(status_code=404, detail="Penjualan tidak ditemukan")
    if doc.get("cashier") != user["email"] and not is_super_admin(user):
        raise HTTPException(status_code=403, detail="Tidak dapat mengubah transaksi admin lain")
    update = payload.dict()
    if not can_see_financials(user):
        # Non-super users never receive modal; keep the real cost and recompute
        # margin from the (editable) selling price so records stay accurate.
        modal = float(doc.get("harga_modal") or 0)
        update["harga_modal"] = modal
        update["margin"] = float(update.get("harga_jual") or 0) - modal
    await db.sales.update_one({"id": sale_id}, {"$set": update})
    updated = await db.sales.find_one({"id": sale_id})
    await log_activity(user, "update", "sale", f"Ubah penjualan {updated.get('nama_barang') or ''}".strip(), sale_id)
    return strip_sale_financials(sale_public(updated), user)

class SaleStatusUpdate(BaseModel):
    sudah_diambil: Optional[str] = None
    metode_pengambilan: Optional[str] = None

@api_router.patch("/admin/sales/{sale_id}/status")
async def update_sale_status(sale_id: str, payload: SaleStatusUpdate, user: dict = Depends(get_current_user)):
    doc = await db.sales.find_one({"id": sale_id, "deleted_at": None})
    if not doc:
        raise HTTPException(status_code=404, detail="Penjualan tidak ditemukan")
    if doc.get("cashier") != user["email"] and not is_super_admin(user):
        raise HTTPException(status_code=403, detail="Tidak dapat mengubah transaksi admin lain")
    updates = {}
    if payload.sudah_diambil is not None:
        if payload.sudah_diambil not in ("Sudah", "Belum", ""):
            raise HTTPException(status_code=400, detail="Status pengambilan tidak valid")
        updates["sudah_diambil"] = payload.sudah_diambil or None
    if payload.metode_pengambilan is not None:
        updates["metode_pengambilan"] = payload.metode_pengambilan or None
    if updates:
        await db.sales.update_one({"id": sale_id}, {"$set": updates})
    updated = await db.sales.find_one({"id": sale_id})
    await log_activity(user, "update", "sale", f"Status pengambilan: {updated.get('sudah_diambil') or 'Belum'} • {updated.get('nama_barang') or ''}".strip(), sale_id)
    return sale_public(updated)

@api_router.delete("/admin/sales/{sale_id}")
async def delete_sale(sale_id: str, user: dict = Depends(get_current_user)):
    doc = await db.sales.find_one({"id": sale_id, "deleted_at": None})
    if not doc:
        raise HTTPException(status_code=404, detail="Penjualan tidak ditemukan")
    if doc.get("cashier") != user["email"] and not is_super_admin(user):
        raise HTTPException(status_code=403, detail="Tidak dapat menghapus transaksi admin lain")
    await db.sales.update_one({"id": sale_id}, {"$set": {"deleted_at": datetime.now(timezone.utc).isoformat()}})
    await log_activity(user, "delete", "sale", f"Hapus penjualan {doc.get('nama_barang') or ''}".strip(), sale_id)
    return {"ok": True}

class SaleReturn(BaseModel):
    returned: bool

@api_router.patch("/admin/sales/{sale_id}/return")
async def return_sale(sale_id: str, payload: SaleReturn, user: dict = Depends(get_current_user)):
    doc = await db.sales.find_one({"id": sale_id, "deleted_at": None})
    if not doc:
        raise HTTPException(status_code=404, detail="Penjualan tidak ditemukan")
    if doc.get("cashier") != user["email"] and not is_super_admin(user):
        raise HTTPException(status_code=403, detail="Tidak dapat mengubah transaksi admin lain")
    returned = bool(payload.returned)
    await db.sales.update_one({"id": sale_id}, {"$set": {"returned": returned, "returned_at": datetime.now(timezone.utc).isoformat() if returned else None}})
    updated = await db.sales.find_one({"id": sale_id})
    await log_activity(user, "update", "sale", f"{'Retur' if returned else 'Batal retur'} penjualan {doc.get('nama_barang') or ''}".strip(), sale_id)
    return strip_sale_financials(sale_public(updated), user)

# ---------------- Service / Bengkel routes ----------------
async def _sync_service_sale(service: dict, cashier_email: str):
    """Sparepart price flows into sales history & omset via a linked sale record."""
    sid = service["id"]
    harga = _num(service, "harga_sparepart")
    existing = await db.sales.find_one({"service_id": sid, "deleted_at": None})
    if harga > 0:
        sale_fields = {
            "tanggal_penjualan": str(service.get("created_at", ""))[:10],
            "nama_pembeli": service.get("nama_customer"),
            "nama_barang": f"Sparepart Servis - {service.get('sepeda', '') or ''}".strip().rstrip("-").strip(),
            "kode_barang": service.get("nomor_servis"),
            "harga_modal": 0,
            "harga_jual": harga,
            "margin": harga,
            "metode_pembayaran": "Servis",
        }
        if existing:
            await db.sales.update_one({"id": existing["id"]}, {"$set": sale_fields})
        else:
            doc = {f: None for f in SALE_FIELDS}
            doc.update(sale_fields)
            doc["id"] = str(uuid.uuid4())
            doc["created_at"] = datetime.now(timezone.utc).isoformat()
            doc["deleted_at"] = None
            doc["cashier"] = cashier_email
            doc["service_id"] = sid
            doc["source"] = "servis"
            await db.sales.insert_one(doc)
    elif existing:
        await db.sales.update_one({"id": existing["id"]}, {"$set": {"deleted_at": datetime.now(timezone.utc).isoformat()}})

@api_router.get("/admin/services")
async def list_services(user: dict = Depends(get_current_user)):
    q = {"deleted_at": None}
    if not is_super_admin(user):
        q["owner_email"] = user["email"]
    docs = await db.services.find(q).sort("created_at", -1).to_list(2000)
    return [service_public(d) for d in docs]

@api_router.post("/admin/services")
async def create_service(payload: ServiceCreate, user: dict = Depends(get_current_user)):
    if payload.status and payload.status not in SERVICE_STATUSES:
        raise HTTPException(status_code=400, detail="Status servis tidak valid")
    now = datetime.now(timezone.utc).isoformat()
    doc = payload.dict()
    doc["id"] = str(uuid.uuid4())
    doc["created_at"] = now
    doc["deleted_at"] = None
    doc["owner_email"] = user.get("email", "")
    if not (doc.get("nomor_servis") or "").strip():
        count = await db.services.count_documents({})
        doc["nomor_servis"] = f"SVC-{datetime.now(timezone.utc).strftime('%y%m%d')}-{count + 1:03d}"
    doc["harga_sparepart"] = _num(doc, "harga_sparepart")
    doc["biaya_jasa"] = _num(doc, "biaya_jasa")
    await db.services.insert_one(doc)
    await _sync_service_sale(doc, user.get("email", ""))
    await log_activity(user, "create", "service", f"Servis {doc.get('nomor_servis')} • {doc.get('nama_customer') or ''}".strip(), doc["id"])
    return service_public(doc)

@api_router.put("/admin/services/{service_id}")
async def update_service(service_id: str, payload: ServiceCreate, user: dict = Depends(get_current_user)):
    doc = await db.services.find_one({"id": service_id, "deleted_at": None})
    if not doc:
        raise HTTPException(status_code=404, detail="Servis tidak ditemukan")
    if doc.get("owner_email") != user["email"] and not is_super_admin(user):
        raise HTTPException(status_code=403, detail="Tidak dapat mengubah servis admin lain")
    if payload.status and payload.status not in SERVICE_STATUSES:
        raise HTTPException(status_code=400, detail="Status servis tidak valid")
    upd = payload.dict()
    upd["harga_sparepart"] = _num(upd, "harga_sparepart")
    upd["biaya_jasa"] = _num(upd, "biaya_jasa")
    if not (upd.get("nomor_servis") or "").strip():
        upd["nomor_servis"] = doc.get("nomor_servis")
    await db.services.update_one({"id": service_id}, {"$set": upd})
    merged = {**doc, **upd}
    await _sync_service_sale(merged, doc.get("owner_email", user["email"]))
    await log_activity(user, "update", "service", f"Ubah servis {merged.get('nomor_servis')} • status {merged.get('status')}", service_id)
    return service_public(merged)

@api_router.delete("/admin/services/{service_id}")
async def delete_service(service_id: str, user: dict = Depends(get_current_user)):
    doc = await db.services.find_one({"id": service_id, "deleted_at": None})
    if not doc:
        raise HTTPException(status_code=404, detail="Servis tidak ditemukan")
    if doc.get("owner_email") != user["email"] and not is_super_admin(user):
        raise HTTPException(status_code=403, detail="Tidak dapat menghapus servis admin lain")
    now = datetime.now(timezone.utc).isoformat()
    await db.services.update_one({"id": service_id}, {"$set": {"deleted_at": now}})
    linked = await db.sales.find_one({"service_id": service_id, "deleted_at": None})
    if linked:
        await db.sales.update_one({"id": linked["id"]}, {"$set": {"deleted_at": now}})
    await log_activity(user, "delete", "service", f"Hapus servis {doc.get('nomor_servis')}", service_id)
    return {"ok": True}

@api_router.get("/admin/services/summary")
async def services_summary(user: dict = Depends(get_current_user)):
    q = {"deleted_at": None}
    if not is_super_admin(user):
        q["owner_email"] = user["email"]
    docs = await db.services.find(q).to_list(5000)
    per_mek = {}
    for d in docs:
        t = ((d.get("teknisi") or "").strip()) or "Tanpa Teknisi"
        m = per_mek.setdefault(t, {"teknisi": t, "jumlah_servis": 0, "total_jasa": 0.0, "total_sparepart": 0.0})
        m["jumlah_servis"] += 1
        m["total_jasa"] += _num(d, "biaya_jasa")
        m["total_sparepart"] += _num(d, "harga_sparepart")
    active = sum(1 for d in docs if d.get("status") not in ("Selesai", "Diambil"))
    if not can_see_financials(user):
        return {
            "total_services": len(docs),
            "active_services": active,
            "total_jasa": None,
            "total_sparepart": None,
            "mechanics": [],
        }
    return {
        "total_services": len(docs),
        "active_services": active,
        "total_jasa": sum(_num(d, "biaya_jasa") for d in docs),
        "total_sparepart": sum(_num(d, "harga_sparepart") for d in docs),
        "mechanics": sorted(per_mek.values(), key=lambda x: -x["total_jasa"]),
    }

# ---------------- POS (Point of Sale) routes ----------------
@api_router.post("/admin/pos/checkout")
async def pos_checkout(payload: PosCheckout, user: dict = Depends(get_current_user)):
    if not payload.items:
        raise HTTPException(status_code=400, detail="Keranjang kosong")
    # Authoritative modal (cost) lookup server-side so margin stays correct
    # even when the cashier/admin UI never receives cost_price.
    cost_map = {}
    obj_ids = []
    for it in payload.items:
        if it.product_id:
            try:
                obj_ids.append(ObjectId(it.product_id))
            except Exception:
                pass
    prod_map = {}
    if obj_ids:
        async for pd in db.products.find({"_id": {"$in": obj_ids}}):
            prod_map[str(pd["_id"])] = pd
            cost_map[str(pd["_id"])] = float(pd.get("cost_price") or 0)
    norm = []
    subtotal = 0.0
    for it in payload.items:
        qty = max(1, int(it.qty or 1))
        unit = max(0.0, float(it.unit_price or 0))
        # Server-authoritative cost: prefer the chosen size's modal price.
        pd = prod_map.get(it.product_id)
        if pd and pd.get("sizes"):
            so = next((s for s in pd["sizes"] if s.get("label") == it.size), None)
            cost = float((so or {}).get("cost_price") if so else pd.get("cost_price") or 0) or 0.0
        elif pd:
            cost = float(pd.get("cost_price") or 0)
        else:
            cost = max(0.0, float(it.cost_price or 0))
        line = unit * qty
        subtotal += line
        norm.append({"product_id": it.product_id, "name": it.name, "code": it.code or "", "color": (it.color or None), "size": (it.size or None), "qty": qty, "unit_price": unit, "cost_price": cost, "line_total": line})
    discount = min(max(0.0, float(payload.discount or 0)), subtotal)
    total = subtotal - discount
    amount_paid = float(payload.amount_paid or 0)
    change = amount_paid - total
    factor = (total / subtotal) if subtotal > 0 else 1.0
    tx_id = str(uuid.uuid4())
    now_dt = datetime.now(timezone.utc)
    now = now_dt.isoformat()
    tanggal_label = id_date_label(now_dt)
    cashier = user.get("email", "")

    for n in norm:
        line_final = round(n["line_total"] * factor)
        cost_total = n["cost_price"] * n["qty"]
        doc = {f: None for f in SALE_FIELDS}
        doc.update({
            "tanggal_penjualan": tanggal_label,
            "nama_pembeli": payload.customer,
            "nama_barang": n["name"],
            "kode_barang": n["code"],
            "ukuran_warna": (f"{n['size']} • " if n.get("size") else "") + (f"{n['color']} • " if n.get("color") else "") + f"{n['qty']} x {rupiah_id(n['unit_price'])}",
            "harga_modal": cost_total,
            "harga_jual": line_final,
            "margin": line_final - cost_total,
            "metode_pembayaran": payload.payment_method,
            "sudah_diambil": payload.taken_status or "Sudah",
            "alamat_pengiriman": payload.address,
            "metode_pengambilan": payload.courier_name,
            "bukti_transfer": payload.transfer_proof,
        })
        doc["id"] = str(uuid.uuid4())
        doc["created_at"] = now
        doc["deleted_at"] = None
        doc["cashier"] = cashier
        doc["source"] = "pos"
        doc["transaction_id"] = tx_id
        doc["qty"] = n["qty"]
        doc["color"] = n.get("color")
        doc["size"] = n.get("size")
        doc["telepon_pembeli"] = payload.customer_phone
        doc["nama_kurir"] = payload.courier_name
        doc["telepon_kurir"] = payload.courier_phone
        await db.sales.insert_one(doc)
        if n["product_id"]:
            try:
                await deduct_stock(n["product_id"], n["qty"], n.get("color"), n.get("size"))
            except Exception as e:
                logger.error(f"Stock update failed for {n['product_id']}: {e}")

    tx_doc = {
        "id": tx_id, "created_at": now, "cashier": cashier, "customer": payload.customer,
        "payment_method": payload.payment_method, "amount_paid": amount_paid, "discount": discount,
        "subtotal": subtotal, "total": total, "change": change, "tanggal": tanggal_label,
        "customer_phone": payload.customer_phone, "address": payload.address,
        "taken_status": payload.taken_status or "Sudah", "courier_name": payload.courier_name,
        "courier_phone": payload.courier_phone, "transfer_proof": payload.transfer_proof,
        "product_note": payload.product_note,
        "items": [{k: n[k] for k in ("name", "code", "color", "size", "qty", "unit_price", "line_total")} for n in norm],
    }
    await db.pos_transactions.insert_one(dict(tx_doc))
    await log_activity(user, "create", "sale", f"Transaksi POS {len(norm)} barang • {rupiah_id(total)}", tx_id)
    return tx_doc

@api_router.post("/admin/products")
async def create_product(payload: ProductCreate, user: dict = Depends(forbid_cashier)):
    if payload.category not in await active_category_names():
        raise HTTPException(status_code=400, detail="Kategori tidak valid")
    doc = payload.model_dump()
    doc["specs"] = payload.specs.model_dump()
    if payload.variants is not None:
        variants = normalize_variants(payload.variants)
    else:
        variants = [{"color": "Standar", "stock": int(doc.get("stock") or 0)}]
    doc["variants"] = variants
    doc["stock"] = sum(v["stock"] for v in variants)
    sizes = normalize_sizes(payload.sizes) if payload.sizes is not None else []
    doc["sizes"] = sizes
    if sizes:
        doc["price"] = min(s["price"] for s in sizes)
        cheapest = min(sizes, key=lambda s: s["price"])
        doc["cost_price"] = cheapest["cost_price"]
        agg = aggregate_variants_from_sizes(sizes)
        doc["variants"] = agg
        doc["stock"] = sum(v["stock"] for v in agg)
    now = datetime.now(timezone.utc).isoformat()
    doc["created_at"] = now
    doc["updated_at"] = now
    res = await db.products.insert_one(doc)
    saved = await db.products.find_one({"_id": res.inserted_id})
    await log_activity(user, "create", "product", f"Tambah produk {doc.get('name') or ''}".strip(), str(res.inserted_id))
    return product_admin(saved, user)

@api_router.put("/admin/products/{product_id}")
async def update_product(product_id: str, payload: ProductUpdate, user: dict = Depends(forbid_cashier)):
    existing = await db.products.find_one({"_id": ObjectId(product_id)})
    if not existing:
        raise HTTPException(status_code=404, detail="Produk tidak ditemukan")
    updates = {k: v for k, v in payload.model_dump(exclude_unset=True).items() if v is not None}
    if "specs" in updates and payload.specs is not None:
        updates["specs"] = payload.specs.model_dump()
    if "variants" in updates and payload.variants is not None:
        variants = normalize_variants(payload.variants)
        updates["variants"] = variants
        updates["stock"] = sum(v["stock"] for v in variants)
    if "sizes" in updates and payload.sizes is not None:
        sizes = normalize_sizes(payload.sizes)
        # Non-super admins can't see/edit modal (cost) price — preserve existing
        # cost per size (matched by label) so their edits never wipe cost data.
        if sizes and not is_super_admin(user):
            existing_cost = {s.get("label"): s.get("cost_price", 0) for s in (existing.get("sizes") or [])}
            for s in sizes:
                if s["label"] in existing_cost:
                    s["cost_price"] = existing_cost[s["label"]]
        updates["sizes"] = sizes
        if sizes:
            updates["price"] = min(s["price"] for s in sizes)
            updates["cost_price"] = min(sizes, key=lambda s: s["price"])["cost_price"]
            agg = aggregate_variants_from_sizes(sizes)
            updates["variants"] = agg
            updates["stock"] = sum(v["stock"] for v in agg)
    if "category" in updates and updates["category"] != existing.get("category") and updates["category"] not in await active_category_names():
        raise HTTPException(status_code=400, detail="Kategori tidak valid")
    updates["updated_at"] = datetime.now(timezone.utc).isoformat()
    # Catat riwayat perubahan harga
    if "price" in updates:
        old_price = int(existing.get("price") or 0)
        new_price = int(updates["price"] or 0)
        if old_price != new_price:
            await db.price_history.insert_one({
                "product_id": str(product_id),
                "product_name": existing.get("name"),
                "old_price": old_price,
                "new_price": new_price,
                "changed_by": user.get("email"),
                "changed_at": datetime.now(timezone.utc).isoformat(),
            })
    res = await db.products.update_one({"_id": ObjectId(product_id)}, {"$set": updates})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Produk tidak ditemukan")
    saved = await db.products.find_one({"_id": ObjectId(product_id)})
    await log_activity(user, "update", "product", f"Ubah produk {saved.get('name') or ''}".strip(), product_id)
    return product_admin(saved, user)

class VariantsUpdate(BaseModel):
    variants: List[Variant]

@api_router.put("/admin/products/{product_id}/variants")
async def update_product_variants(product_id: str, payload: VariantsUpdate, user: dict = Depends(forbid_cashier)):
    existing = await db.products.find_one({"_id": ObjectId(product_id)})
    if not existing:
        raise HTTPException(status_code=404, detail="Produk tidak ditemukan")
    variants = normalize_variants(payload.variants)
    total = sum(v["stock"] for v in variants)
    await db.products.update_one({"_id": ObjectId(product_id)}, {"$set": {"variants": variants, "stock": total, "updated_at": datetime.now(timezone.utc).isoformat()}})
    saved = await db.products.find_one({"_id": ObjectId(product_id)})
    return product_admin(saved, user)

async def deduct_stock(product_id: str, qty: int, color: Optional[str] = None, size: Optional[str] = None):
    prod = await db.products.find_one({"_id": ObjectId(product_id)})
    if not prod:
        return
    sizes = prod.get("sizes") or []
    if sizes:
        # Nested product: deduct from the specific size -> color, then rebuild
        # the aggregate variants + total from the authoritative nested data.
        size_obj = None
        if size:
            size_obj = next((s for s in sizes if s.get("label") == size), None)
        if size_obj is None:
            size_obj = max(sizes, key=lambda s: sum(int(c.get("stock", 0)) for c in (s.get("colors") or [])))
        cols = size_obj.get("colors") or []
        if cols:
            target = None
            if color:
                target = next((c for c in cols if c.get("color") == color), None)
            if target is None:
                target = max(cols, key=lambda c: int(c.get("stock", 0)))
            target["stock"] = max(0, int(target.get("stock", 0)) - qty)
        agg = aggregate_variants_from_sizes(sizes)
        total = sum(v["stock"] for v in agg)
        await db.products.update_one({"_id": ObjectId(product_id)}, {"$set": {"sizes": sizes, "variants": agg, "stock": total}})
        return
    variants = prod.get("variants") or []
    if variants:
        target = None
        if color:
            target = next((v for v in variants if v.get("color") == color), None)
        if target is None:
            target = max(variants, key=lambda v: int(v.get("stock", 0)))
        target["stock"] = max(0, int(target.get("stock", 0)) - qty)
        total = sum(int(v.get("stock", 0)) for v in variants)
        await db.products.update_one({"_id": ObjectId(product_id)}, {"$set": {"variants": variants, "stock": total}})
    else:
        await db.products.update_one(
            {"_id": ObjectId(product_id)},
            [{"$set": {"stock": {"$max": [0, {"$subtract": [{"$ifNull": ["$stock", 0]}, qty]}]}}}],
        )

def price_history_out(doc):
    return {
        "id": str(doc.get("_id")),
        "product_id": doc.get("product_id"),
        "product_name": doc.get("product_name"),
        "old_price": doc.get("old_price"),
        "new_price": doc.get("new_price"),
        "changed_by": doc.get("changed_by"),
        "changed_at": doc.get("changed_at"),
    }

@api_router.get("/admin/price-history")
async def price_history_all(user: dict = Depends(get_current_user)):
    docs = await db.price_history.find().sort("changed_at", -1).to_list(200)
    return [price_history_out(d) for d in docs]

@api_router.get("/admin/products/{product_id}/price-history")
async def price_history_product(product_id: str, user: dict = Depends(get_current_user)):
    docs = await db.price_history.find({"product_id": str(product_id)}).sort("changed_at", -1).to_list(100)
    return [price_history_out(d) for d in docs]

@api_router.delete("/admin/products/{product_id}")
async def delete_product(product_id: str, user: dict = Depends(forbid_cashier)):
    existing = await db.products.find_one({"_id": ObjectId(product_id)})
    res = await db.products.delete_one({"_id": ObjectId(product_id)})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Produk tidak ditemukan")
    await log_activity(user, "delete", "product", f"Hapus produk {(existing or {}).get('name') or ''}".strip(), product_id)
    return {"message": "Produk dihapus"}

@api_router.post("/admin/upload")
async def upload_image(file: UploadFile = File(...), user: dict = Depends(get_current_user)):
    ext = file.filename.split(".")[-1].lower() if "." in file.filename else "png"
    content_type = MIME_TYPES.get(ext, file.content_type or "image/png")
    path = f"{APP_NAME}/products/{uuid.uuid4()}.{ext}"
    data = await file.read()
    result = put_object(path, data, content_type)
    stored_path = result["path"]
    await db.files.insert_one({
        "id": str(uuid.uuid4()), "storage_path": stored_path, "original_filename": file.filename,
        "content_type": content_type, "size": result.get("size", len(data)),
        "is_deleted": False, "created_at": datetime.now(timezone.utc).isoformat()
    })
    return {"image_url": f"/api/products/image/{stored_path}"}

# ---------------- Seed ----------------
SEED_PRODUCTS = [
    {"name": "SK Trail Pro 29", "category": "Sepeda Gunung", "price": 8500000, "stock": 12, "status": "Tersedia",
     "description": "Sepeda gunung hardtail agresif untuk trail teknikal dengan geometri modern dan handling presisi.",
     "image_url": "https://images.unsplash.com/photo-1621122940876-2b3be129159c?crop=entropy&cs=srgb&fm=jpg&ixid=M3w3NTY2NjZ8MHwxfHNlYXJjaHwzfHxtb3VudGFpbiUyMGJpa2UlMjBiaWN5Y2xlJTIwYWN0aW9uJTIwY3ljbGluZ3xlbnwwfHx8fDE3ODkxOTYxMjJ8MA&ixlib=rb-4.1.0&q=85",
     "specs": {"frame": "Aluminium Alloy 6061", "transmisi": "Shimano Deore 12-Speed", "rem": "Hydraulic Disc Brake", "ukuran_roda": "29 inci", "baterai_motor": "-"}},
    {"name": "SK Enduro X1", "category": "Sepeda Gunung", "price": 15750000, "stock": 5, "status": "Stok Terbatas",
     "description": "Full suspension enduro untuk medan ekstrem, travel 160mm dan performa turunan luar biasa.",
     "image_url": "https://images.unsplash.com/photo-1606087492572-424ebe0f2f61?crop=entropy&cs=srgb&fm=jpg&ixid=M3w3NTY2NjZ8MHwxfHNlYXJjaHwyfHxtb3VudGFpbiUyMGJpa2UlMjBiaWN5Y2xlJTIwYWN0aW9uJTIwY3ljbGluZ3xlbnwwfHx8fDE3ODkxOTYxMjJ8MA&ixlib=rb-4.1.0&q=85",
     "specs": {"frame": "Carbon Fiber", "transmisi": "SRAM GX Eagle 12-Speed", "rem": "SRAM Code R 4-Piston", "ukuran_roda": "29 inci", "baterai_motor": "-"}},
    {"name": "SK Street BMX 20", "category": "BMX", "price": 3200000, "stock": 20, "status": "Tersedia",
     "description": "BMX freestyle tangguh untuk trik park dan street, ringan namun kokoh.",
     "image_url": "https://images.unsplash.com/photo-1594015197353-373334c20904?crop=entropy&cs=srgb&fm=jpg&ixid=M3w3NTY2NzB8MHwxfHNlYXJjaHwyfHxCTVglMjByaWRlciUyMHRyaWNrfGVufDB8fHx8MTc4NzcyMzIyOHww&ixlib=rb-4.1.0&q=85",
     "specs": {"frame": "Chromoly 4130", "transmisi": "Single Speed 25T", "rem": "U-Brake Rear", "ukuran_roda": "20 inci", "baterai_motor": "-"}},
    {"name": "SK Pro Park BMX", "category": "BMX", "price": 4100000, "stock": 8, "status": "Tersedia",
     "description": "BMX kompetisi dengan komponen premium untuk rider tingkat lanjut.",
     "image_url": "https://images.unsplash.com/flagged/photo-1553677969-1d67bbe9d55a?crop=entropy&cs=srgb&fm=jpg&ixid=M3w3NTY2NzB8MHwxfHNlYXJjaHwxfHxCTVglMjByaWRlciUyMHRyaWNrfGVufDB8fHx8MTc4NzcyMzIyOHww&ixlib=rb-4.1.0&q=85",
     "specs": {"frame": "Full Chromoly", "transmisi": "Single Speed 25/9", "rem": "Gyro U-Brake", "ukuran_roda": "20 inci", "baterai_motor": "-"}},
    {"name": "SK Junior 16", "category": "Sepeda Anak", "price": 1450000, "stock": 25, "status": "Tersedia",
     "description": "Sepeda anak aman dan ceria dengan roda bantu, cocok usia 4-7 tahun.",
     "image_url": "https://images.unsplash.com/photo-1507035895480-2b3156c31fc8?crop=entropy&cs=srgb&fm=jpg&fit=max&q=80&w=1080",
     "specs": {"frame": "Hi-Ten Steel", "transmisi": "Single Speed", "rem": "Coaster + V-Brake", "ukuran_roda": "16 inci", "baterai_motor": "-"}},
    {"name": "SK Kids Explorer 20", "category": "Sepeda Anak", "price": 1850000, "stock": 0, "status": "Inden",
     "description": "Sepeda anak petualang dengan 6-speed untuk usia 8-11 tahun.",
     "image_url": "https://images.unsplash.com/photo-1507035895480-2b3156c31fc8?crop=entropy&cs=srgb&fm=jpg&fit=max&q=80&w=1080",
     "specs": {"frame": "Aluminium Alloy", "transmisi": "Shimano 6-Speed", "rem": "V-Brake", "ukuran_roda": "20 inci", "baterai_motor": "-"}},
    {"name": "SK Fold Urban 20", "category": "Sepeda Lipat", "price": 4750000, "stock": 15, "status": "Tersedia",
     "description": "Sepeda lipat urban ringkas, dilipat dalam 15 detik untuk mobilitas kota.",
     "image_url": "https://images.unsplash.com/photo-1485965120184-e220f721d03e?crop=entropy&cs=srgb&fm=jpg&fit=max&q=80&w=1080",
     "specs": {"frame": "Aluminium Alloy Lipat", "transmisi": "Shimano 7-Speed", "rem": "Dual V-Brake", "ukuran_roda": "20 inci", "baterai_motor": "-"}},
    {"name": "SK Fold Lite Carbon", "category": "Sepeda Lipat", "price": 9800000, "stock": 4, "status": "Stok Terbatas",
     "description": "Sepeda lipat carbon ultra ringan hanya 9.8kg, ideal untuk komuter premium.",
     "image_url": "https://images.unsplash.com/photo-1485965120184-e220f721d03e?crop=entropy&cs=srgb&fm=jpg&fit=max&q=80&w=1080",
     "specs": {"frame": "Carbon Fiber Lipat", "transmisi": "Shimano 9-Speed", "rem": "Hydraulic Disc", "ukuran_roda": "20 inci", "baterai_motor": "-"}},
    {"name": "SK E-Trail 750W", "category": "Sepeda Listrik", "price": 22500000, "stock": 7, "status": "Tersedia",
     "description": "E-bike gunung bertenaga motor 750W, jangkauan hingga 80km per pengisian.",
     "image_url": "https://images.unsplash.com/photo-1571068316344-75bc76f77890?crop=entropy&cs=srgb&fm=jpg&fit=max&q=80&w=1080",
     "specs": {"frame": "Aluminium Alloy E-Bike", "transmisi": "Shimano 9-Speed", "rem": "Hydraulic Disc 180mm", "ukuran_roda": "27.5 inci", "baterai_motor": "Motor 750W / Baterai 48V 15Ah"}},
    {"name": "SK E-City Cruiser", "category": "Sepeda Listrik", "price": 16900000, "stock": 10, "status": "Tersedia",
     "description": "E-bike perkotaan nyaman dengan desain elegan dan bantuan pedal cerdas.",
     "image_url": "https://images.unsplash.com/photo-1532298229144-0ec0c57515c7?crop=entropy&cs=srgb&fm=jpg&fit=max&q=80&w=1080",
     "specs": {"frame": "Aluminium Step-Through", "transmisi": "Shimano 7-Speed", "rem": "Tektro Disc Brake", "ukuran_roda": "700c", "baterai_motor": "Motor 500W / Baterai 36V 12Ah"}},
]

EXTRA_ADMINS = []

async def seed_admin():
    admin_email = os.environ["ADMIN_EMAIL"].lower()
    admin_password = os.environ["ADMIN_PASSWORD"]
    existing = await db.users.find_one({"email": admin_email})
    if existing is None:
        await db.users.insert_one({"email": admin_email, "password_hash": hash_password(admin_password), "name": "Admin SK Bike", "role": "admin", "created_at": datetime.now(timezone.utc).isoformat()})
        logger.info("Admin user seeded")
    elif not verify_password(admin_password, existing["password_hash"]):
        await db.users.update_one({"email": admin_email}, {"$set": {"password_hash": hash_password(admin_password)}})
    for extra in EXTRA_ADMINS:
        email = extra["email"].lower()
        found = await db.users.find_one({"email": email})
        if found is None:
            await db.users.insert_one({"email": email, "password_hash": hash_password(extra["password"]), "name": extra.get("name", "Admin"), "role": "admin", "created_at": datetime.now(timezone.utc).isoformat()})
            logger.info(f"Extra admin seeded: {email}")

async def seed_products():
    count = await db.products.count_documents({})
    if count == 0:
        now = datetime.now(timezone.utc).isoformat()
        for p in SEED_PRODUCTS:
            p = {**p, "created_at": now, "updated_at": now}
            await db.products.insert_one(p)
        logger.info("Products seeded")

async def migrate_variants():
    async for doc in db.products.find({"variants": {"$exists": False}}):
        stock = int(doc.get("stock") or 0)
        await db.products.update_one({"_id": doc["_id"]}, {"$set": {"variants": [{"color": "Standar", "stock": stock}]}})

@app.on_event("startup")
async def startup():
    await db.users.create_index("email", unique=True)
    await db.login_attempts.create_index("identifier")
    await seed_admin()
    await seed_products()
    await seed_categories()
    await migrate_variants()
    try:
        init_storage()
        logger.info("Storage initialized")
    except Exception as e:
        logger.error(f"Storage init failed: {e}")

@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()

app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=[os.environ.get("FRONTEND_URL", "http://localhost:3000"), "http://localhost:3000"],
    allow_origin_regex=".*",
    allow_methods=["*"],
    allow_headers=["*"],
)
